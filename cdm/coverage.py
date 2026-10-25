"""SI² Coverage model: classes and subclasses, coverage catalog, terms, endorsements and synthetic policies.
One catalog generates the coverage CDM (workbook, SQL Server, CSV) and the coverage ontology (Turtle, JSON-LD).

    python coverage.py OUT_DIR

Personal lines: HO (homeowners), PA (personal auto). Commercial lines: CP, GL, CA, WC, and the commercial package (CPP).
All codes, limits, rates and rules are illustrative. Forms and codes belong to ISO, NCCI, AAIS, ACORD and the state bureaus.
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

import pandas as pd
from rdflib import RDF, RDFS, OWL, XSD, SKOS, DCTERMS, BNode, Graph, Literal, Namespace, URIRef

VERSION = "2026.10"
KEY = ["country", "state", "company", "line", "lob", "product", "object", "module", "effective_from", "effective_to"]
LEVELS = [("Policy", "Applies to the whole policy term"), ("Line", "Applies to one line or coverage part in the policy"),
          ("Coverage", "Attaches to another coverage and follows it"), ("Object", "Attaches to one insured object"),
          ("Endorsement", "Changes coverage at any of the other levels")]
LOBS = {"HO": ("Personal", "Homeowners"), "PA": ("Personal", "Personal auto"), "CP": ("Commercial", "Commercial property"),
        "GL": ("Commercial", "General liability"), "CA": ("Commercial", "Commercial auto"), "WC": ("Commercial", "Workers' compensation")}

# ------------------------------------------------------------------------------------------ class hierarchy
# class, parent, kind (Coverage/Term/Endorsement), definition
CLASSES = [
    ("Coverage", None, "Coverage", "A promise to pay or defend for a defined loss, under stated terms."),
    ("FirstPartyCoverage", "Coverage", "Coverage", "Pays the insured for its own loss."),
    ("PropertyCoverage", "FirstPartyCoverage", "Coverage", "Direct physical loss to property."),
    ("DwellingCoverage", "PropertyCoverage", "Coverage", "The home and structures on a residence premises."),
    ("PersonalPropertyCoverage", "PropertyCoverage", "Coverage", "Contents and belongings of a household."),
    ("CommercialBuildingCoverage", "PropertyCoverage", "Coverage", "A commercial building."),
    ("BusinessPersonalPropertyCoverage", "PropertyCoverage", "Coverage", "Contents, stock and equipment of a business."),
    ("AutoPhysicalDamageCoverage", "PropertyCoverage", "Coverage", "Damage to a covered vehicle."),
    ("CollisionCoverage", "AutoPhysicalDamageCoverage", "Coverage", "Collision with another object or upset."),
    ("ComprehensiveCoverage", "AutoPhysicalDamageCoverage", "Coverage", "Other than collision: theft, fire, hail, glass."),
    ("TimeElementCoverage", "FirstPartyCoverage", "Coverage", "Loss that grows with time after property damage."),
    ("LossOfUseCoverage", "TimeElementCoverage", "Coverage", "Additional living expense and fair rental value."),
    ("BusinessIncomeCoverage", "TimeElementCoverage", "Coverage", "Lost income and extra expense of a business."),
    ("FirstPartyInjuryCoverage", "FirstPartyCoverage", "Coverage", "Medical and income benefits regardless of fault."),
    ("PersonalInjuryProtection", "FirstPartyInjuryCoverage", "Coverage", "No-fault benefits required in some states."),
    ("MedicalPaymentsCoverage", "FirstPartyInjuryCoverage", "Coverage", "Medical expenses of others or occupants, regardless of fault."),
    ("LiabilityCoverage", "Coverage", "Coverage", "Pays others for injury or damage the insured is legally liable for, and defends."),
    ("BodilyInjuryPropertyDamageLiability", "LiabilityCoverage", "Coverage", "Bodily injury and property damage to others."),
    ("AutoLiability", "BodilyInjuryPropertyDamageLiability", "Coverage", "Liability from owning or using autos."),
    ("PremisesOperationsLiability", "BodilyInjuryPropertyDamageLiability", "Coverage", "Liability from premises and operations."),
    ("PersonalLiability", "BodilyInjuryPropertyDamageLiability", "Coverage", "Liability of a household."),
    ("ProductsCompletedOperationsLiability", "LiabilityCoverage", "Coverage", "Liability from products and completed work."),
    ("PersonalAdvertisingInjuryLiability", "LiabilityCoverage", "Coverage", "Libel, slander, false arrest, advertising injury."),
    ("EmployersLiability", "LiabilityCoverage", "Coverage", "Employee injury claims outside workers' compensation."),
    ("UninsuredMotoristsCoverage", "Coverage", "Coverage", "Injury caused by an uninsured or underinsured driver."),
    ("StatutoryBenefitsCoverage", "Coverage", "Coverage", "Benefits set by statute."),
    ("WorkersCompensationBenefits", "StatutoryBenefitsCoverage", "Coverage", "Benefits under the workers' compensation law of a state."),
    ("CoverageTerm", None, "Term", "A value that defines how much or how a coverage pays."),
    ("Limit", "CoverageTerm", "Term", "Most the coverage pays."),
    ("PerOccurrenceLimit", "Limit", "Term", "Most paid for one occurrence."),
    ("AggregateLimit", "Limit", "Term", "Most paid in the policy term."),
    ("PerPersonLimit", "Limit", "Term", "Most paid for one person."),
    ("CombinedSingleLimit", "Limit", "Term", "One limit for bodily injury and property damage."),
    ("Sublimit", "Limit", "Term", "A smaller limit inside a parent limit."),
    ("Deductible", "CoverageTerm", "Term", "Amount the insured bears before the coverage pays."),
    ("FlatDeductible", "Deductible", "Term", "A fixed amount."),
    ("PercentDeductible", "Deductible", "Term", "A percentage of a stated value."),
    ("HurricaneDeductible", "PercentDeductible", "Term", "Percentage deductible for a hurricane or named storm."),
    ("WindHailDeductible", "PercentDeductible", "Term", "Percentage deductible for windstorm or hail."),
    ("EarthquakeDeductible", "PercentDeductible", "Term", "Percentage deductible for earthquake."),
    ("Coinsurance", "CoverageTerm", "Term", "Share of value the insured must carry."),
    ("Valuation", "CoverageTerm", "Term", "How a loss is valued."),
    ("WaitingPeriod", "CoverageTerm", "Term", "Time before a time-element coverage starts."),
    ("Endorsement", None, "Endorsement", "A form that changes the policy."),
    ("AddCoverageEndorsement", "Endorsement", "Endorsement", "Adds a coverage."),
    ("ModifyTermEndorsement", "Endorsement", "Endorsement", "Changes a limit, deductible or other term."),
    ("BroadenEndorsement", "Endorsement", "Endorsement", "Widens what a coverage covers or who is insured."),
    ("RestrictEndorsement", "Endorsement", "Endorsement", "Narrows a coverage or its valuation."),
    ("ExclusionEndorsement", "Endorsement", "Endorsement", "Removes a cause of loss, object or activity."),
    ("StateAmendatoryEndorsement", "Endorsement", "Endorsement", "Changes conditions to meet a state's law."),
]

# ------------------------------------------------------------------------------------------ coverage catalog
# code, name, lob, class, attach_level, attaches_to_object, parent_coverage
COVERAGES = [
    ("HO_A", "Coverage A - Dwelling", "HO", "DwellingCoverage", "Object", "Dwelling", None),
    ("HO_B", "Coverage B - Other Structures", "HO", "DwellingCoverage", "Coverage", None, "HO_A"),
    ("HO_C", "Coverage C - Personal Property", "HO", "PersonalPropertyCoverage", "Coverage", None, "HO_A"),
    ("HO_D", "Coverage D - Loss of Use", "HO", "LossOfUseCoverage", "Coverage", None, "HO_A"),
    ("HO_E", "Coverage E - Personal Liability", "HO", "PersonalLiability", "Line", None, None),
    ("HO_F", "Coverage F - Medical Payments to Others", "HO", "MedicalPaymentsCoverage", "Line", None, None),
    ("HO_SPP", "Scheduled Personal Property", "HO", "PersonalPropertyCoverage", "Policy", None, None),
    ("HO_EQ", "Earthquake", "HO", "DwellingCoverage", "Coverage", None, "HO_A"),
    ("PA_BI", "Bodily Injury Liability", "PA", "AutoLiability", "Line", None, None),
    ("PA_PD", "Property Damage Liability", "PA", "AutoLiability", "Line", None, None),
    ("PA_UM", "Uninsured Motorists", "PA", "UninsuredMotoristsCoverage", "Line", None, None),
    ("PA_PIP", "Personal Injury Protection", "PA", "PersonalInjuryProtection", "Line", None, None),
    ("PA_MED", "Medical Payments", "PA", "MedicalPaymentsCoverage", "Line", None, None),
    ("PA_COLL", "Collision", "PA", "CollisionCoverage", "Object", "Vehicle", None),
    ("PA_COMP", "Comprehensive", "PA", "ComprehensiveCoverage", "Object", "Vehicle", None),
    ("CP_BLDG", "Building", "CP", "CommercialBuildingCoverage", "Object", "Building", None),
    ("CP_BPP", "Business Personal Property", "CP", "BusinessPersonalPropertyCoverage", "Object", "Building", None),
    ("CP_BI", "Business Income and Extra Expense", "CP", "BusinessIncomeCoverage", "Coverage", None, "CP_BLDG"),
    ("GL_OCC", "Coverage A - Bodily Injury and Property Damage", "GL", "PremisesOperationsLiability", "Line", None, None),
    ("GL_PAI", "Coverage B - Personal and Advertising Injury", "GL", "PersonalAdvertisingInjuryLiability", "Line", None, None),
    ("GL_MED", "Coverage C - Medical Payments", "GL", "MedicalPaymentsCoverage", "Line", None, None),
    ("GL_PCO", "Products-Completed Operations", "GL", "ProductsCompletedOperationsLiability", "Line", None, None),
    ("CA_LIAB", "Auto Liability", "CA", "AutoLiability", "Line", None, None),
    ("CA_UM", "Uninsured Motorists", "CA", "UninsuredMotoristsCoverage", "Line", None, None),
    ("CA_PIP", "Personal Injury Protection", "CA", "PersonalInjuryProtection", "Line", None, None),
    ("CA_COLL", "Collision", "CA", "CollisionCoverage", "Object", "Vehicle", None),
    ("CA_COMP", "Comprehensive", "CA", "ComprehensiveCoverage", "Object", "Vehicle", None),
    ("WC_PART1", "Part One - Workers' Compensation", "WC", "WorkersCompensationBenefits", "Line", None, None),
    ("WC_EL", "Part Two - Employers Liability", "WC", "EmployersLiability", "Line", None, None),
]
# coverage, term_code, term class, data type, unit, pick list, rule (how the value is set)
TERMS = [
    ("HO_A", "LIMIT", "Limit", "AMOUNT", "USD", None, "Insured to replacement cost"),
    ("HO_A", "VALUATION", "Valuation", "CODE", None, "RC|ACV", "Default RC"),
    ("HO_A", "AOP_DED", "FlatDeductible", "AMOUNT", "USD", "500|1000|2500|5000", "Default 1000"),
    ("HO_A", "HURR_DED", "HurricaneDeductible", "PERCENT", "% of HO_A", "2|5|10", "Policy level where mandatory"),
    ("HO_B", "LIMIT", "Sublimit", "AMOUNT", "USD", None, "10% of HO_A"),
    ("HO_C", "LIMIT", "Sublimit", "AMOUNT", "USD", None, "50% of HO_A"),
    ("HO_D", "LIMIT", "Sublimit", "AMOUNT", "USD", None, "20% of HO_A"),
    ("HO_E", "LIMIT", "PerOccurrenceLimit", "AMOUNT", "USD", "100000|300000|500000", "Default 300000"),
    ("HO_F", "LIMIT", "PerPersonLimit", "AMOUNT", "USD", "1000|5000", "Default 5000"),
    ("HO_SPP", "LIMIT", "Limit", "AMOUNT", "USD", None, "Sum of scheduled items"),
    ("HO_EQ", "EQ_DED", "EarthquakeDeductible", "PERCENT", "% of HO_A", "10|15|20", "Default 15"),
    ("PA_BI", "PER_PERSON", "PerPersonLimit", "AMOUNT", "USD", "25000|50000|100000|250000", None),
    ("PA_BI", "PER_ACCIDENT", "PerOccurrenceLimit", "AMOUNT", "USD", "50000|100000|300000|500000", None),
    ("PA_PD", "LIMIT", "PerOccurrenceLimit", "AMOUNT", "USD", "25000|50000|100000", None),
    ("PA_UM", "PER_PERSON", "PerPersonLimit", "AMOUNT", "USD", None, "Up to BI limits"),
    ("PA_PIP", "LIMIT", "PerPersonLimit", "AMOUNT", "USD", None, "State minimum"),
    ("PA_MED", "LIMIT", "PerPersonLimit", "AMOUNT", "USD", "1000|5000|10000", None),
    ("PA_COLL", "DED", "FlatDeductible", "AMOUNT", "USD", "250|500|1000", "Default 500"),
    ("PA_COMP", "DED", "FlatDeductible", "AMOUNT", "USD", "100|250|500|1000", "Default 250"),
    ("CP_BLDG", "LIMIT", "Limit", "AMOUNT", "USD", None, "Insured value"),
    ("CP_BLDG", "COINS", "Coinsurance", "CODE", None, "80|90|100", "Default 90"),
    ("CP_BLDG", "VALUATION", "Valuation", "CODE", None, "RC|ACV|AGREED", "Default RC"),
    ("CP_BLDG", "AOP_DED", "FlatDeductible", "AMOUNT", "USD", "1000|5000|10000|25000", "Default 5000"),
    ("CP_BLDG", "WIND_DED", "WindHailDeductible", "PERCENT", "% of building value", "1|2|5", "Where mandatory"),
    ("CP_BPP", "LIMIT", "Limit", "AMOUNT", "USD", None, "Insured value"),
    ("CP_BI", "LIMIT", "Limit", "AMOUNT", "USD", None, "12 months of income"),
    ("CP_BI", "WAITING", "WaitingPeriod", "CODE", "hours", "24|72", "Default 72"),
    ("GL_OCC", "EACH_OCC", "PerOccurrenceLimit", "AMOUNT", "USD", "500000|1000000|2000000", "Default 1000000"),
    ("GL_OCC", "GEN_AGG", "AggregateLimit", "AMOUNT", "USD", None, "2 x each occurrence"),
    ("GL_PAI", "LIMIT", "PerPersonLimit", "AMOUNT", "USD", None, "Equal to each occurrence"),
    ("GL_MED", "LIMIT", "PerPersonLimit", "AMOUNT", "USD", "5000|10000", "Default 5000"),
    ("GL_PCO", "AGG", "AggregateLimit", "AMOUNT", "USD", None, "Equal to general aggregate"),
    ("CA_LIAB", "CSL", "CombinedSingleLimit", "AMOUNT", "USD", "500000|1000000|2000000", "Default 1000000"),
    ("CA_UM", "LIMIT", "CombinedSingleLimit", "AMOUNT", "USD", None, "Up to CSL"),
    ("CA_PIP", "LIMIT", "PerPersonLimit", "AMOUNT", "USD", None, "State minimum"),
    ("CA_COLL", "DED", "FlatDeductible", "AMOUNT", "USD", "500|1000|2500", "Default 1000"),
    ("CA_COMP", "DED", "FlatDeductible", "AMOUNT", "USD", "500|1000|2500", "Default 1000"),
    ("WC_PART1", "LIMIT", "Limit", "CODE", None, "STATUTORY", "Statutory"),
    ("WC_EL", "EACH_ACC", "PerOccurrenceLimit", "AMOUNT", "USD", "100000|500000|1000000", "Default 1000000"),
    ("WC_EL", "DIS_POL", "AggregateLimit", "AMOUNT", "USD", None, "Equal to each accident"),
    ("WC_EL", "DIS_EE", "PerPersonLimit", "AMOUNT", "USD", None, "Equal to each accident"),
]
# applicability: coverage[.term], scope, offer_status, default, reason
APPLIC = [
    ("HO_A", {"lob": "HO", "product": "HO-3 2024"}, "Mandatory", None, "Base coverage of the HO-3 form."),
    ("HO_B", {"lob": "HO"}, "Mandatory", "10% of HO_A", "Follows Coverage A."),
    ("HO_C", {"lob": "HO"}, "Mandatory", "50% of HO_A", "Follows Coverage A."),
    ("HO_D", {"lob": "HO"}, "Mandatory", "20% of HO_A", "Follows Coverage A."),
    ("HO_E", {"lob": "HO"}, "Mandatory", "300000", "Section II liability."),
    ("HO_F", {"lob": "HO"}, "Mandatory", "5000", "Section II medical payments."),
    ("HO_SPP", {"lob": "HO"}, "Offered", None, "By endorsement."),
    ("HO_EQ", {"lob": "HO"}, "Not available", None, "Not offered countrywide in this illustration."),
    ("HO_EQ", {"lob": "HO", "state": "CA"}, "Offered", None, "Earthquake must be offered in California."),
    ("HO_A.HURR_DED", {"lob": "HO"}, "Not available", None, "No hurricane deductible outside listed states."),
    ("HO_A.HURR_DED", {"lob": "HO", "state": "FL"}, "Mandatory", "2", "Florida hurricane deductible, policy level."),
    ("HO_A.HURR_DED", {"lob": "HO", "state": "TX"}, "Mandatory", "2", "Coastal windstorm deductible."),
    ("PA_BI", {"lob": "PA"}, "Mandatory", "25000/50000", "Financial responsibility law."),
    ("PA_PD", {"lob": "PA"}, "Mandatory", "25000", "Financial responsibility law."),
    ("PA_UM", {"lob": "PA"}, "Offered", None, "Rejection must be signed where allowed."),
    ("PA_PIP", {"lob": "PA"}, "Not available", None, "Not a no-fault state."),
    ("PA_PIP", {"lob": "PA", "state": "FL"}, "Mandatory", "10000", "Florida no-fault."),
    ("PA_PIP", {"lob": "PA", "state": "NY"}, "Mandatory", "50000", "New York no-fault."),
    ("PA_PIP", {"lob": "PA", "state": "MI"}, "Mandatory", None, "Michigan no-fault; levels chosen by the insured."),
    ("PA_PIP", {"lob": "PA", "state": "NJ"}, "Mandatory", None, "New Jersey no-fault."),
    ("PA_COLL", {"lob": "PA"}, "Offered", "500", "Optional physical damage."),
    ("PA_COMP", {"lob": "PA"}, "Offered", "250", "Optional physical damage."),
    ("CP_BLDG", {"lob": "CP"}, "Offered", None, "Building or contents, at least one."),
    ("CP_BPP", {"lob": "CP"}, "Offered", None, "Building or contents, at least one."),
    ("CP_BI", {"lob": "CP"}, "Offered", None, "Time element."),
    ("CP_BLDG.WIND_DED", {"lob": "CP"}, "Offered", None, "Optional elsewhere."),
    ("CP_BLDG.WIND_DED", {"lob": "CP", "state": "TX"}, "Mandatory", "2", "Texas coastal wind and hail."),
    ("CP_BLDG.WIND_DED", {"lob": "CP", "state": "FL"}, "Mandatory", "2", "Florida named storm."),
    ("GL_OCC", {"lob": "GL"}, "Mandatory", "1000000", "Coverage A."),
    ("GL_PAI", {"lob": "GL"}, "Mandatory", None, "Coverage B."),
    ("GL_MED", {"lob": "GL"}, "Offered", "5000", "Coverage C."),
    ("GL_PCO", {"lob": "GL"}, "Mandatory", None, "Products-completed operations aggregate."),
    ("CA_LIAB", {"lob": "CA"}, "Mandatory", "1000000", "Auto liability."),
    ("CA_UM", {"lob": "CA"}, "Offered", None, "Rejection must be signed where allowed."),
    ("CA_PIP", {"lob": "CA"}, "Not available", None, "Not a no-fault state."),
    ("CA_PIP", {"lob": "CA", "state": "FL"}, "Mandatory", "10000", "Florida no-fault."),
    ("CA_PIP", {"lob": "CA", "state": "NY"}, "Mandatory", "50000", "New York no-fault."),
    ("CA_COLL", {"lob": "CA"}, "Offered", "1000", "Optional physical damage."),
    ("CA_COMP", {"lob": "CA"}, "Offered", "1000", "Optional physical damage."),
    ("WC_PART1", {"lob": "WC"}, "Mandatory", "STATUTORY", "Statutory benefits."),
    ("WC_PART1", {"lob": "WC", "state": "OH"}, "Not available", None, "Monopolistic state fund."),
    ("WC_PART1", {"lob": "WC", "state": "WA"}, "Not available", None, "Monopolistic state fund."),
    ("WC_PART1", {"lob": "WC", "state": "ND"}, "Not available", None, "Monopolistic state fund."),
    ("WC_PART1", {"lob": "WC", "state": "WY"}, "Not available", None, "Monopolistic state fund."),
    ("WC_EL", {"lob": "WC"}, "Mandatory", "1000000", "Part Two."),
    ("WC_EL", {"lob": "WC", "state": "OH"}, "Offered", "1000000", "Stop-gap employers liability only."),
]
# endorsement code, title, class, attach level, lob, scope, effects [(effect, target coverage, target term, value)]
ENDORSEMENTS = [
    ("END_HO_SPP", "Scheduled personal property", "AddCoverageEndorsement", "Policy", "HO", {},
     [("Add", "HO_SPP", None, "schedule")]),
    ("END_HO_CINC", "Increased Coverage C", "ModifyTermEndorsement", "Coverage", "HO", {},
     [("Modify term", "HO_C", "LIMIT", "70% of HO_A")]),
    ("END_HO_ROOF_ACV", "Roof surfacing at actual cash value", "RestrictEndorsement", "Coverage", "HO", {},
     [("Restrict", "HO_A", "VALUATION", "ACV for roof surfacing")]),
    ("END_HO_EQ", "Earthquake", "AddCoverageEndorsement", "Coverage", "HO", {"state": "CA"},
     [("Add", "HO_EQ", None, None), ("Modify term", "HO_EQ", "EQ_DED", "15")]),
    ("END_HO_FL_AMEND", "Special provisions - Florida", "StateAmendatoryEndorsement", "Policy", "HO", {"state": "FL"},
     [("Modify condition", None, None, "Cancellation and notice per state law")]),
    ("END_CP_WIND_EXCL", "Windstorm or hail exclusion", "ExclusionEndorsement", "Object", "CP", {},
     [("Exclude", "CP_BLDG", None, "Windstorm or hail")]),
    ("END_GL_AI_OWNER", "Additional insured - owner (scheduled)", "BroadenEndorsement", "Coverage", "GL", {},
     [("Broaden", "GL_OCC", None, "Scheduled owner as additional insured")]),
    ("END_GL_DESIG_OPS", "Designated operations exclusion", "ExclusionEndorsement", "Line", "GL", {},
     [("Exclude", "GL_OCC", None, "Scheduled operations")]),
    ("END_CA_HNOA", "Hired and non-owned auto", "AddCoverageEndorsement", "Line", "CA", {},
     [("Add", "CA_LIAB", None, "Symbols 8 and 9")]),
    ("END_CA_TX_AMEND", "Texas changes", "StateAmendatoryEndorsement", "Line", "CA", {"state": "TX"},
     [("Modify condition", None, None, "Texas conditions")]),
    ("END_WC_OTHER_STATES", "Other states insurance (3.C)", "BroadenEndorsement", "Line", "WC", {},
     [("Broaden", "WC_PART1", None, "States listed in 3.C")]),
]
PACKAGES = [  # product, lob, required
    ("HO-3 2024", "HO", "Required"), ("PA-Std 2024", "PA", "Required"),
    ("CPP 2025", "CP", "Optional"), ("CPP 2025", "GL", "Optional"), ("CPP 2025", "CA", "Optional"),
    ("CPP 2025", "IM", "Optional"), ("CP-Std 2025", "CP", "Required"), ("GL-Std 2025", "GL", "Required"),
    ("WC-Std 2025", "WC", "Required"),
]

COV = {c[0]: c for c in COVERAGES}


def full_scope(scope: dict, lob: str | None = None) -> dict:
    row = {k: "*" for k in KEY}
    row.update({"effective_from": "2025-01-01", "effective_to": "9999-12-31"})
    if lob:
        row["lob"] = lob
        row["line"] = LOBS[lob][0]
    row.update(scope)
    if row["lob"] != "*":
        row["line"] = LOBS[row["lob"]][0]
    return row


def resolve(target: str, ctx: dict):
    best, bs = None, -1
    for t, s, st, d, why in APPLIC:
        if t == target and all(ctx.get(k) == v for k, v in s.items()) and len(s) > bs:
            best, bs = (st, d, why), len(s)
    return best


# ------------------------------------------------------------------------------------------ catalog tables
def catalog() -> dict[str, pd.DataFrame]:
    t = {}
    depth = {}
    parents = {c: p for c, p, _k, _d in CLASSES}

    def d(c):
        return 0 if parents[c] is None else 1 + d(parents[c])
    t["Levels"] = pd.DataFrame(LEVELS, columns=["attach_level", "meaning"])
    t["Classes"] = pd.DataFrame([{"class": c, "parent_class": p or "", "kind": k, "depth": d(c), "definition": de,
                                  "path": " > ".join(reversed([x for x in iter_path(c, parents)]))} for c, p, k, de in CLASSES])
    t["CoverageDef"] = pd.DataFrame([{"coverage_code": c, "name": n, "line": LOBS[l][0], "lob": l, "class": k,
                                      "class_path": " > ".join(reversed(list(iter_path(k, parents)))), "attach_level": a,
                                      "attaches_to_object": o or "", "parent_coverage_code": p or ""}
                                     for c, n, l, k, a, o, p in COVERAGES])
    t["CoverageTermDef"] = pd.DataFrame([{"coverage_code": c, "term_code": tc, "term_class": k, "data_type": dt,
                                          "unit": u or "", "options": o or "", "rule": r or ""} for c, tc, k, dt, u, o, r in TERMS])
    ap = []
    for i, (target, scope, st, default, why) in enumerate(APPLIC, 1):
        cov, _, term = target.partition(".")
        ap.append({"applicability_id": f"CA{i:03d}", "coverage_code": cov, "term_code": term, **full_scope(scope),
                   "offer_status": st, "default_value": default or "", "specificity": len(scope), "reason": why})
    t["CoverageApplicability"] = pd.DataFrame(ap)
    t["EndorsementDef"] = pd.DataFrame([{"endorsement_code": c, "title": ti, "class": k, "attach_level": a,
                                         **full_scope(s, lob)} for c, ti, k, a, lob, s, _e in ENDORSEMENTS])
    t["EndorsementEffect"] = pd.DataFrame([{"endorsement_code": c, "seq": j, "effect_type": e, "target_coverage_code": tc or "",
                                            "target_term_code": tt or "", "value": v or ""}
                                           for c, _ti, _k, _a, _l, _s, effs in ENDORSEMENTS for j, (e, tc, tt, v) in enumerate(effs, 1)])
    t["PackageDef"] = pd.DataFrame(PACKAGES, columns=["product_code", "lob", "required"])
    return t


def iter_path(c, parents):
    while c:
        yield c
        c = parents[c]


# ------------------------------------------------------------------------------------------ synthetic policies
def pct(base, p):
    return round(base * p / 100)


def synth() -> dict[str, pd.DataFrame]:
    pol, line, obj, cov, term, end, prem = ([] for _ in range(7))
    n = {"ln": 0, "ob": 0, "cv": 0, "en": 0}

    def P(pno, product, state, company, insured, eff):
        pol.append({"policy_number": pno, "product_code": product, "country": "US", "state": state, "company": company,
                    "named_insured": insured, "effective_date": eff, "expiration_date": str(int(eff[:4]) + 1) + eff[4:]})

    def L(pno, lob):
        n["ln"] += 1
        lid = f"LN{n['ln']:03d}"
        line.append({"line_id": lid, "policy_number": pno, "lob": lob, "line": LOBS[lob][0]})
        return lid

    def O(pno, lid, otype, desc, value=None):
        n["ob"] += 1
        oid = f"OB{n['ob']:03d}"
        obj.append({"object_id": oid, "policy_number": pno, "line_id": lid, "object_type": otype, "description": desc, "value": value})
        return oid

    def C(pno, code, attach_to, terms: dict, premium=None, state=None, parent=None, source="Base form"):
        n["cv"] += 1
        cid = f"CV{n['cv']:03d}"
        c = COV[code]
        cov.append({"coverage_id": cid, "policy_number": pno, "coverage_code": code, "name": c[1], "class": c[3],
                    "attach_level": c[4], "attach_to_id": attach_to, "parent_coverage_id": parent or "", "source": source})
        for tc, (v, src) in terms.items():
            term.append({"coverage_id": cid, "term_code": tc, "value": v, "source": src})
        if premium:
            prem.append({"coverage_id": cid, "policy_number": pno, "state": state, "transaction": "NB", "written_premium": premium})
        return cid

    def E(pno, code, attach_to, note=""):
        n["en"] += 1
        e = next(x for x in ENDORSEMENTS if x[0] == code)
        end.append({"endorsement_id": f"EN{n['en']:03d}", "policy_number": pno, "endorsement_code": code, "title": e[1],
                    "class": e[2], "attach_level": e[3], "attach_to_id": attach_to, "note": note})

    def homeowners(pno, state, a, insured, prem_a, hurr=None, eq=False, endorsements=()):
        P(pno, "HO-3 2024", state, "SI2-ADM", insured, "2026-03-01")
        lid = L(pno, "HO")
        dw = O(pno, lid, "Dwelling", f"Single-family dwelling, {state}", a)
        terms = {"LIMIT": (a, "Selected"), "VALUATION": ("RC", "Default"), "AOP_DED": (1000, "Default")}
        ca = C(pno, "HO_A", dw, terms, prem_a, state)
        c_pct = 70 if "END_HO_CINC" in endorsements else 50
        C(pno, "HO_B", ca, {"LIMIT": (pct(a, 10), "Rule 10% of HO_A")}, parent=ca)
        C(pno, "HO_C", ca, {"LIMIT": (pct(a, c_pct), f"{'Endorsement' if c_pct == 70 else 'Rule'} {c_pct}% of HO_A")}, parent=ca)
        C(pno, "HO_D", ca, {"LIMIT": (pct(a, 20), "Rule 20% of HO_A")}, parent=ca)
        C(pno, "HO_E", lid, {"LIMIT": (300000, "Default")}, round(prem_a * 0.06), state)
        C(pno, "HO_F", lid, {"LIMIT": (5000, "Default")})
        if hurr:
            # policy-level hurricane deductible: recorded as a term on Coverage A, attached at Policy level
            term.append({"coverage_id": ca, "term_code": "HURR_DED", "value": hurr, "source": f"Mandatory in {state}; applies at Policy level"})
        if eq:
            C(pno, "HO_EQ", ca, {"EQ_DED": (15, "Endorsement")}, round(prem_a * 0.4), state, parent=ca, source="END_HO_EQ")
            E(pno, "END_HO_EQ", ca)
        for e in endorsements:
            if e == "END_HO_SPP":
                C(pno, "HO_SPP", pno, {"LIMIT": (42000, "Schedule: ring, watch, painting")}, 380, state, source="END_HO_SPP")
                E(pno, e, pno, "3 scheduled items")
            elif e in ("END_HO_CINC", "END_HO_ROOF_ACV"):
                E(pno, e, ca)
            else:
                E(pno, e, pno)

    homeowners("HO2026001", "FL", 450000, "Rivera Household", 6120, hurr="2%", endorsements=("END_HO_SPP", "END_HO_ROOF_ACV", "END_HO_FL_AMEND"))
    homeowners("HO2026002", "CA", 820000, "Nguyen Household", 2950, eq=True, endorsements=("END_HO_CINC",))
    homeowners("HO2026003", "OH", 310000, "Baker Household", 1410)

    # personal auto, New York (PIP)
    P("PA2026001", "PA-Std 2024", "NY", "SI2-ADM", "Shah Household", "2026-05-10")
    lid = L("PA2026001", "PA")
    C("PA2026001", "PA_BI", lid, {"PER_PERSON": (100000, "Selected"), "PER_ACCIDENT": (300000, "Selected")}, 820, "NY")
    C("PA2026001", "PA_PD", lid, {"LIMIT": (100000, "Selected")}, 410, "NY")
    C("PA2026001", "PA_PIP", lid, {"LIMIT": (50000, "Mandatory in NY")}, 290, "NY")
    C("PA2026001", "PA_UM", lid, {"PER_PERSON": (100000, "Selected")}, 140, "NY")
    for vin in ("2023 sedan", "2021 SUV"):
        v = O("PA2026001", lid, "Vehicle", vin, None)
        C("PA2026001", "PA_COLL", v, {"DED": (500, "Default")}, 520, "NY")
        C("PA2026001", "PA_COMP", v, {"DED": (250, "Default")}, 190, "NY")

    # commercial package, Texas: CP + GL + CA
    P("CPP2026001", "CPP 2025", "TX", "SI2-ADM", "Lone Star Fabrication", "2026-04-01")
    cp = L("CPP2026001", "CP")
    b = O("CPP2026001", cp, "Building", "Plant, Houston", 6600000)
    bl = C("CPP2026001", "CP_BLDG", b, {"LIMIT": (6600000, "Selected"), "COINS": (90, "Default"), "VALUATION": ("RC", "Default"),
                                         "AOP_DED": (5000, "Default"), "WIND_DED": ("2%", "Mandatory in TX")}, 38400, "TX")
    C("CPP2026001", "CP_BPP", b, {"LIMIT": (1500000, "Selected")}, 7900, "TX")
    C("CPP2026001", "CP_BI", bl, {"LIMIT": (1200000, "Selected"), "WAITING": (72, "Default")}, 5100, "TX", parent=bl)
    gl = L("CPP2026001", "GL")
    occ = C("CPP2026001", "GL_OCC", gl, {"EACH_OCC": (1000000, "Default"), "GEN_AGG": (2000000, "Rule 2 x each occurrence")}, 21800, "TX")
    C("CPP2026001", "GL_PAI", gl, {"LIMIT": (1000000, "Rule")})
    C("CPP2026001", "GL_MED", gl, {"LIMIT": (5000, "Default")})
    C("CPP2026001", "GL_PCO", gl, {"AGG": (2000000, "Rule")}, 6400, "TX")
    E("CPP2026001", "END_GL_AI_OWNER", occ, "Owner: Port Holdings LLC (synthetic)")
    ca_ = L("CPP2026001", "CA")
    C("CPP2026001", "CA_LIAB", ca_, {"CSL": (1000000, "Default")}, 9800, "TX")
    C("CPP2026001", "CA_UM", ca_, {"LIMIT": (30000, "Selected")}, 600, "TX")
    for desc, val in (("Heavy truck VEH0005", 140000), ("Light truck VEH0006", 52000)):
        v = O("CPP2026001", ca_, "Vehicle", desc, val)
        C("CPP2026001", "CA_COLL", v, {"DED": (1000, "Default")}, 2100 if val > 100000 else 900, "TX")
        C("CPP2026001", "CA_COMP", v, {"DED": (1000, "Default")}, 700 if val > 100000 else 310, "TX")
    E("CPP2026001", "END_CA_HNOA", ca_)
    E("CPP2026001", "END_CA_TX_AMEND", ca_)

    # monoline CP Florida with wind exclusion (wind written elsewhere)
    P("CP2026001", "CP-Std 2025", "FL", "SI2-ES", "Harbor Point Holdings", "2026-06-15")
    cp = L("CP2026001", "CP")
    b = O("CP2026001", cp, "Building", "Office, Miami", 8100000)
    C("CP2026001", "CP_BLDG", b, {"LIMIT": (8100000, "Selected"), "COINS": (90, "Default"), "VALUATION": ("RC", "Default"),
                                   "AOP_DED": (10000, "Selected")}, 31200, "FL")
    E("CP2026001", "END_CP_WIND_EXCL", b, "Wind placed with a wind pool (synthetic)")

    # GL New York contractor
    P("GL2026001", "GL-Std 2025", "NY", "SI2-ADM", "Empire Ridge Builders", "2026-02-01")
    gl = L("GL2026001", "GL")
    O("GL2026001", gl, "Classification", "Carpentry contractor, payroll 5,300,000", 5300000)
    occ = C("GL2026001", "GL_OCC", gl, {"EACH_OCC": (1000000, "Default"), "GEN_AGG": (2000000, "Rule")}, 186000, "NY")
    C("GL2026001", "GL_PAI", gl, {"LIMIT": (1000000, "Rule")})
    C("GL2026001", "GL_PCO", gl, {"AGG": (2000000, "Rule")}, 31300, "NY")
    E("GL2026001", "END_GL_DESIG_OPS", gl, "Excludes roofing work (synthetic)")

    # WC Pennsylvania with Ohio on 3.C
    P("WC2026001", "WC-Std 2025", "PA", "SI2-ADM", "Keystone Plumbing", "2026-01-01")
    wc = L("WC2026001", "WC")
    O("WC2026001", wc, "Classification", "Plumbing 5183, payroll 1,250,000", 1250000)
    C("WC2026001", "WC_PART1", wc, {"LIMIT": ("STATUTORY", "Statutory")}, 61200, "PA")
    C("WC2026001", "WC_EL", wc, {"EACH_ACC": (1000000, "Default"), "DIS_POL": (1000000, "Rule"), "DIS_EE": (1000000, "Rule")})
    E("WC2026001", "END_WC_OTHER_STATES", wc, "3.C states listed; Ohio excluded (monopolistic) - stop-gap EL only")

    df = pd.DataFrame
    return {"Policy": df(pol), "PolicyLine": df(line), "PolicyObject": df(obj), "PolicyCoverage": df(cov),
            "PolicyCoverageTerm": df(term), "PolicyEndorsement": df(end), "PolicyCoveragePremium": df(prem)}


def in_force(data: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Coverage in force: each coverage with its level, what it attaches to, its terms and the endorsements acting on it."""
    cov, term, end = data["PolicyCoverage"], data["PolicyCoverageTerm"], data["PolicyEndorsement"]
    objs = data["PolicyObject"].set_index("object_id")
    lines = data["PolicyLine"].set_index("line_id")
    covs = cov.set_index("coverage_id")

    def label(x):
        if x in objs.index:
            return f"Object {x}: {objs.at[x, 'description']}"
        if x in lines.index:
            return f"Line {x}: {lines.at[x, 'lob']}"
        if x in covs.index:
            return f"Coverage {x}: {covs.at[x, 'coverage_code']}"
        return f"Policy {x}"
    def line_of(x):
        if x in lines.index:
            return x
        if x in objs.index:
            return objs.at[x, "line_id"]
        if x in covs.index:
            return line_of(covs.at[x, "attach_to_id"])
        return None
    rows = []
    for c in cov.to_dict("records"):
        c = pd.Series(c)
        t = term[term.coverage_id == c.coverage_id]
        acting = []
        for e in end.itertuples():
            if e.policy_number != c.policy_number:
                continue
            effs = next(x for x in ENDORSEMENTS if x[0] == e.endorsement_code)[6]
            targets = {tc for _e, tc, _t, _v in effs if tc}
            if c.coverage_code in targets:
                acting.append(e.endorsement_code)
            elif not targets and (e.attach_level == "Policy" or line_of(e.attach_to_id) == line_of(c.coverage_id)):
                acting.append(e.endorsement_code)
        rows.append({"policy_number": c.policy_number, "coverage_id": c.coverage_id, "coverage_code": c.coverage_code, "name": c.name,
                     "class": c["class"], "attach_level": c.attach_level,
                     "attaches_to": label(c.attach_to_id),
                     "terms": "; ".join(f"{r.term_code}={r.value}" for r in t.itertuples()),
                     "endorsements_acting": ", ".join(acting)})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------------------------------ outputs
def sql_lit(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return str(v)
    return "N'" + str(v).replace("'", "''") + "'"


def sql_type(s: pd.Series) -> str:
    if pd.api.types.is_integer_dtype(s):
        return "BIGINT"
    if pd.api.types.is_float_dtype(s):
        return "DECIMAL(18,2)"
    n = int(s.dropna().astype(str).str.len().max() or 0) if len(s.dropna()) else 0
    for size in (40, 120, 400):
        if n <= size * 0.8:
            return f"NVARCHAR({size})"
    return "NVARCHAR(MAX)"


PK = {"Levels": "attach_level", "Classes": "class", "CoverageDef": "coverage_code", "CoverageApplicability": "applicability_id",
      "EndorsementDef": "endorsement_code", "Policy": "policy_number", "PolicyLine": "line_id", "PolicyObject": "object_id",
      "PolicyCoverage": "coverage_id", "PolicyEndorsement": "endorsement_id"}


def sql_script(tables: dict[str, pd.DataFrame]) -> str:
    out = ["-- SI² Coverage CDM: classes, catalog, terms, endorsements and synthetic policies (schema cov).",
           f"-- Version {VERSION}. Generated from the same catalog as SI2_Coverage_Ontology.ttl. Illustrative.",
           "SET NOCOUNT ON;", "GO", "IF NOT EXISTS (SELECT 1 FROM sys.schemas WHERE name = 'cov') EXEC('CREATE SCHEMA cov');", "GO"]
    for name, df in tables.items():
        df = df.copy()
        # mixed-type value column -> text
        for c in df.columns:
            if df[c].map(type).nunique() > 1 and df[c].dropna().map(lambda x: isinstance(x, str)).any():
                df[c] = df[c].map(lambda x: None if x is None or (isinstance(x, float) and pd.isna(x)) else str(x))
        cols = list(df.columns)
        out.append(f"IF OBJECT_ID('cov.{name}') IS NOT NULL DROP TABLE cov.[{name}];")
        defs = [f"  [{c}] {sql_type(df[c])} {'NOT NULL' if PK.get(name) == c else 'NULL'}" for c in cols]
        if name in PK:
            defs.append(f"  CONSTRAINT PK_cov_{name} PRIMARY KEY ([{PK[name]}])")
        out.append(f"CREATE TABLE cov.[{name}] (\n" + ",\n".join(defs) + "\n);")
        rows = [", ".join(sql_lit(None if (not isinstance(v, str) and pd.isna(v)) else v) for v in r)
                for r in df.astype(object).itertuples(index=False)]
        for i in range(0, len(rows), 500):
            out.append(f"INSERT INTO cov.[{name}] ({', '.join(f'[{c}]' for c in cols)}) VALUES\n(" + "),\n(".join(rows[i:i + 500]) + ");")
        out.append("GO")
    out.append("""-- Which coverages a context may sell: most specific scope wins
CREATE OR ALTER FUNCTION cov.fn_CoverageOffer (@coverage NVARCHAR(40), @term NVARCHAR(40), @state NVARCHAR(10), @lob NVARCHAR(10),
       @product NVARCHAR(60), @company NVARCHAR(20), @asof DATE)
RETURNS TABLE AS RETURN
SELECT TOP 1 a.coverage_code, a.term_code, a.offer_status, a.default_value, a.reason
FROM cov.CoverageApplicability a
WHERE a.coverage_code = @coverage AND a.term_code = ISNULL(@term, N'')
  AND a.state IN ('*', @state) AND a.lob IN ('*', @lob) AND a.product IN ('*', @product) AND a.company IN ('*', @company)
  AND @asof BETWEEN a.effective_from AND a.effective_to
ORDER BY a.specificity DESC;
GO
-- Class hierarchy: every catalog coverage with all its ancestor classes
CREATE OR ALTER VIEW cov.vw_CoverageClassPath AS
WITH p AS (SELECT d.coverage_code, c.class, c.parent_class, 0 AS lvl FROM cov.CoverageDef d JOIN cov.Classes c ON c.class = d.class
           UNION ALL SELECT p.coverage_code, c.class, c.parent_class, p.lvl + 1 FROM p JOIN cov.Classes c ON c.class = p.parent_class)
SELECT coverage_code, class, lvl FROM p;
GO""")
    return "\n".join(out) + "\n"


SI2 = Namespace("https://si2insure.com/ontology/coverage#")
EX = Namespace("https://si2insure.com/data/coverage-sample#")


def ontology(cat: dict[str, pd.DataFrame]) -> Graph:
    g = Graph()
    g.bind("si2c", SI2); g.bind("skos", SKOS); g.bind("owl", OWL); g.bind("dcterms", DCTERMS)
    ont = URIRef(str(SI2).rstrip("#"))
    g.add((ont, RDF.type, OWL.Ontology)); g.add((ont, OWL.versionInfo, Literal(VERSION)))
    g.add((ont, DCTERMS.title, Literal("SI² Coverage Ontology: classes, catalog, terms and endorsements")))
    for c, p, _k, de in CLASSES:
        g.add((SI2[c], RDF.type, OWL.Class)); g.add((SI2[c], RDFS.label, Literal(c)))
        g.add((SI2[c], SKOS.definition, Literal(de)))
        if p:
            g.add((SI2[c], RDFS.subClassOf, SI2[p]))
    for name, label in [("Policy", "Policy"), ("PolicyLine", "Policy line"), ("InsuredObject", "Insured object"),
                        ("Dwelling", "Dwelling"), ("Building", "Building"), ("Vehicle", "Vehicle"), ("Classification", "Classification"),
                        ("EndorsementEffect", "Endorsement effect"), ("CoverageApplicability", "Coverage applicability")]:
        g.add((SI2[name], RDF.type, OWL.Class)); g.add((SI2[name], RDFS.label, Literal(label)))
    for sub in ("Dwelling", "Building", "Vehicle", "Classification"):
        g.add((SI2[sub], RDFS.subClassOf, SI2.InsuredObject))
    # levels
    sch = SI2.AttachLevels
    g.add((sch, RDF.type, SKOS.ConceptScheme)); g.add((sch, RDFS.label, Literal("Coverage attachment levels")))
    for lv, de in LEVELS:
        g.add((SI2["Level_" + lv], RDF.type, SKOS.Concept)); g.add((SI2["Level_" + lv], SKOS.inScheme, sch))
        g.add((SI2["Level_" + lv], SKOS.prefLabel, Literal(lv))); g.add((SI2["Level_" + lv], SKOS.definition, Literal(de)))
    props = [("attachesTo", OWL.ObjectProperty, "attaches to (policy, line, coverage or object)"),
             ("attachLevel", OWL.ObjectProperty, "attach level"), ("attachesToObjectType", OWL.ObjectProperty, "attaches to object type"),
             ("parentCoverage", OWL.ObjectProperty, "follows coverage"), ("hasTerm", OWL.ObjectProperty, "has term"),
             ("hasTermType", OWL.ObjectProperty, "has term type"), ("hasEffect", OWL.ObjectProperty, "has effect"),
             ("targetCoverage", OWL.ObjectProperty, "target coverage"), ("appliesTo", OWL.ObjectProperty, "applies to"),
             ("inLine", OWL.ObjectProperty, "in line"), ("value", OWL.DatatypeProperty, "value"),
             ("termCode", OWL.DatatypeProperty, "term code"), ("rule", OWL.DatatypeProperty, "rule"),
             ("options", OWL.DatatypeProperty, "options"), ("effectType", OWL.DatatypeProperty, "effect type"),
             ("offerStatus", OWL.DatatypeProperty, "offer status"), ("defaultValue", OWL.DatatypeProperty, "default value"),
             ("lob", OWL.DatatypeProperty, "line of business"), ("writtenPremium", OWL.DatatypeProperty, "written premium")]
    for p, t, l in props:
        g.add((SI2[p], RDF.type, t)); g.add((SI2[p], RDFS.label, Literal(l)))
    for k in KEY:
        g.add((SI2["key_" + k], RDF.type, OWL.DatatypeProperty)); g.add((SI2["key_" + k], RDFS.label, Literal(k)))
    # catalog coverages as subclasses of their class
    for c, n, lob, k, a, o, p in COVERAGES:
        u = SI2["Cov_" + c]
        g.add((u, RDF.type, OWL.Class)); g.add((u, RDFS.subClassOf, SI2[k])); g.add((u, RDFS.label, Literal(n)))
        g.add((u, SKOS.notation, Literal(c))); g.add((u, SI2.lob, Literal(lob))); g.add((u, SI2.attachLevel, SI2["Level_" + a]))
        if o:
            g.add((u, SI2.attachesToObjectType, SI2[o]))
        if p:
            g.add((u, SI2.parentCoverage, SI2["Cov_" + p]))
    for c, tc, k, dt, unit, opts, rule in TERMS:
        u = SI2[f"Term_{c}_{tc}"]
        g.add((u, RDF.type, OWL.Class)); g.add((u, RDFS.subClassOf, SI2[k])); g.add((u, SI2.termCode, Literal(tc)))
        g.add((u, RDFS.label, Literal(f"{COV[c][1]}: {tc}")))
        g.add((SI2["Cov_" + c], SI2.hasTermType, u))
        if opts:
            g.add((u, SI2.options, Literal(opts.replace("|", ", "))))
        if rule:
            g.add((u, SI2.rule, Literal(rule)))
    for r in cat["CoverageApplicability"].itertuples(index=False):
        u = SI2[r.applicability_id]
        g.add((u, RDF.type, SI2.CoverageApplicability))
        g.add((u, SI2.appliesTo, SI2[f"Term_{r.coverage_code}_{r.term_code}"] if r.term_code else SI2["Cov_" + r.coverage_code]))
        for k in KEY:
            g.add((u, SI2["key_" + k], Literal(getattr(r, k))))
        g.add((u, SI2.offerStatus, Literal(r.offer_status)))
        if r.default_value:
            g.add((u, SI2.defaultValue, Literal(r.default_value)))
        g.add((u, RDFS.comment, Literal(r.reason)))
    for code, title, k, a, lob, scope, effects in ENDORSEMENTS:
        u = SI2["End_" + code]
        g.add((u, RDF.type, OWL.Class)); g.add((u, RDFS.subClassOf, SI2[k])); g.add((u, RDFS.label, Literal(title)))
        g.add((u, SI2.attachLevel, SI2["Level_" + a])); g.add((u, SI2.lob, Literal(lob)))
        for kk, v in scope.items():
            g.add((u, SI2["key_" + kk], Literal(v)))
        for j, (e, tc, tt, v) in enumerate(effects, 1):
            eff = SI2[f"Eff_{code}_{j}"]
            g.add((u, SI2.hasEffect, eff)); g.add((eff, RDF.type, SI2.EndorsementEffect)); g.add((eff, SI2.effectType, Literal(e)))
            if tc:
                g.add((eff, SI2.targetCoverage, SI2["Cov_" + tc]))
            if tt:
                g.add((eff, SI2.termCode, Literal(tt)))
            if v:
                g.add((eff, SI2.value, Literal(v)))
    return g


def knowledge_graph(data: dict[str, pd.DataFrame]) -> Graph:
    g = Graph()
    g.bind("si2c", SI2); g.bind("ex", EX)
    for r in data["Policy"].itertuples():
        u = EX[r.policy_number]
        g.add((u, RDF.type, SI2.Policy)); g.add((u, SI2.key_state, Literal(r.state))); g.add((u, SI2.key_product, Literal(r.product_code)))
        g.add((u, SI2.key_company, Literal(r.company))); g.add((u, RDFS.label, Literal(r.named_insured)))
    for r in data["PolicyLine"].itertuples():
        g.add((EX[r.line_id], RDF.type, SI2.PolicyLine)); g.add((EX[r.line_id], SI2.lob, Literal(r.lob)))
        g.add((EX[r.line_id], SI2.attachesTo, EX[r.policy_number]))
    for r in data["PolicyObject"].itertuples():
        t = r.object_type if r.object_type in ("Dwelling", "Building", "Vehicle", "Classification") else "InsuredObject"
        g.add((EX[r.object_id], RDF.type, SI2[t])); g.add((EX[r.object_id], RDFS.label, Literal(r.description)))
        g.add((EX[r.object_id], SI2.inLine, EX[r.line_id]))
    for r in data["PolicyCoverage"].itertuples():
        u = EX[r.coverage_id]
        g.add((u, RDF.type, SI2["Cov_" + r.coverage_code])); g.add((u, SI2.attachesTo, EX[r.attach_to_id]))
        g.add((u, SI2.attachLevel, SI2["Level_" + r.attach_level]))
    for i, r in enumerate(data["PolicyCoverageTerm"].itertuples()):
        b = EX[f"{r.coverage_id}_{r.term_code}"]
        g.add((EX[r.coverage_id], SI2.hasTerm, b))
        g.add((b, RDF.type, SI2[f"Term_{data['PolicyCoverage'].set_index('coverage_id').at[r.coverage_id, 'coverage_code']}_{r.term_code}"]))
        g.add((b, SI2.value, Literal(r.value))); g.add((b, RDFS.comment, Literal(r.source)))
    for r in data["PolicyEndorsement"].itertuples():
        g.add((EX[r.endorsement_id], RDF.type, SI2["End_" + r.endorsement_code])); g.add((EX[r.endorsement_id], SI2.attachesTo, EX[r.attach_to_id]))
        g.add((EX[r.endorsement_id], SI2.attachLevel, SI2["Level_" + r.attach_level]))
    for r in data["PolicyCoveragePremium"].itertuples():
        g.add((EX[r.coverage_id], SI2.writtenPremium, Literal(r.written_premium)))
    return g


def write_xlsx(path, sheets):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0B1B2B")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(df.columns, 1):
                w = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
                ws.column_dimensions[get_column_letter(i)].width = 140 if name == "README" else min(60, max(10, w + 2))
            if len(df):
                ws.auto_filter.ref = ws.dimensions


def main(out: Path):
    out.mkdir(parents=True, exist_ok=True)
    cat, data = catalog(), synth()
    data["CoverageInForce"] = in_force(data)
    readme = pd.DataFrame({"SI² Coverage CDM": [
        f"Version {VERSION}. One catalog generates this workbook, the SQL Server script (schema cov), the coverage ontology and a sample knowledge graph.",
        "Attachment levels: Policy, Line, Coverage, Object, plus Endorsement acting on any of them. Every coverage instance stores attach_level + attach_to_id.",
        "Classes: Coverage > FirstPartyCoverage / LiabilityCoverage / UninsuredMotoristsCoverage / StatutoryBenefitsCoverage, down to each catalog coverage; CoverageTerm and Endorsement have their own trees.",
        "Context key on every applicability row: country, state, company, line, lob, product, object, module, effective dates. '*' = all; most specific wins.",
        "Synthetic policies: HO-3 in FL, CA and OH; personal auto NY; commercial package TX (CP + GL + CA); CP FL on E&S paper; GL NY; WC PA with 3.C.",
        "All codes, limits, rates and rules are illustrative; verify against the forms and manuals your company licenses.",
    ]})
    sheets = {"README": readme, **cat, **data}
    write_xlsx(out / "SI2_Coverage_CDM.xlsx", sheets)
    (out / "SI2_Coverage_CDM_SQLServer.sql").write_text(sql_script({**cat, **data}), encoding="utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in cat.items():
            z.writestr(f"catalog/{k}.csv", v.to_csv(index=False))
        for k, v in data.items():
            z.writestr(f"policies/{k}.csv", v.to_csv(index=False))
    (out / "SI2_Coverage_CDM_csv.zip").write_bytes(buf.getvalue())
    g, kg = ontology(cat), knowledge_graph(data)
    (out / "SI2_Coverage_Ontology.ttl").write_text(g.serialize(format="turtle"), encoding="utf-8")
    (out / "SI2_Coverage_Ontology.jsonld").write_text(g.serialize(format="json-ld", indent=1), encoding="utf-8")
    (out / "SI2_Coverage_KnowledgeGraph_Sample.ttl").write_text(kg.serialize(format="turtle"), encoding="utf-8")
    files = ["SI2_Coverage_CDM.xlsx", "SI2_Coverage_CDM_SQLServer.sql", "SI2_Coverage_CDM_csv.zip", "SI2_Coverage_Ontology.ttl",
             "SI2_Coverage_Ontology.jsonld", "SI2_Coverage_KnowledgeGraph_Sample.ttl"]
    with zipfile.ZipFile(out / "SI2_Coverage_CDM_and_Ontology.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for f in files:
            z.write(out / f, f)
    return {"classes": len(CLASSES), "coverages": len(COVERAGES), "terms": len(TERMS), "applicability": len(APPLIC),
            "endorsements": len(ENDORSEMENTS), "effects": len(cat["EndorsementEffect"]), "ontology_triples": len(g), "kg_triples": len(kg),
            **{f"rows_{k}": len(v) for k, v in data.items()}}


if __name__ == "__main__":
    for k, v in main(Path(sys.argv[1] if len(sys.argv) > 1 else "cov_out")).items():
        print(f"{k:28s} {v}")
