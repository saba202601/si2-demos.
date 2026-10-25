"""Generate the SI² Commercial Lines CDM and ontology from ONE store (model.py), plus synthetic data.

    python build.py OUT_DIR

Writes:
  SI2_Commercial_CDM.xlsx            the CDM workbook: keys, objects, elements, applicability, pick lists, rules, rating, sample data
  SI2_Commercial_CDM_SQLServer.sql   SQL Server DDL (schema cdm) + metadata + sample data inserts
  SI2_Commercial_CDM_csv.zip         every table as CSV
  SI2_Commercial_Ontology.ttl        OWL / RDFS / SKOS ontology (Turtle)
  SI2_Commercial_Ontology.jsonld     the same graph as JSON-LD
  SI2_Commercial_CDM_and_Ontology.zip  all of the above
"""
from __future__ import annotations

import io
import random
import sys
import zipfile
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from rdflib import RDF, RDFS, OWL, XSD, SKOS, DCTERMS, Graph, Literal, Namespace, URIRef

import model as M

VERSION = "2026.10"
STAR = "*"
REQ = {"Y": "Required", "N": "Optional", "NA": "Not applicable"}


# ------------------------------------------------------------------------------------------ the store as tables
def scope_row(scope: dict) -> dict:
    row = {k: STAR for k in M.KEY}
    row.update({"line": "Commercial", "effective_from": "2025-01-01", "effective_to": "9999-12-31"})
    row.update({k: v for k, v in scope.items()})
    return row


def tables() -> dict[str, pd.DataFrame]:
    t = {}
    t["Keys"] = pd.DataFrame(M.KEY_DOC, columns=["key", "meaning", "examples", "why it matters"])
    t["Objects"] = pd.DataFrame(M.OBJECTS, columns=["object", "name", "parent_object", "definition", "key_attribute"])
    t["Relations"] = pd.DataFrame(M.RELATIONS, columns=["subject_object", "relation", "object_object", "cardinality"])
    el = pd.DataFrame(M.ELEMENTS)
    el["synonyms"] = el.synonyms.map(" | ".join)
    el["acord_forms"] = el.element_code.map(M.FORM_USE).fillna("")
    el["status"] = "Approved"
    el["approved_by"] = "SME (demo)"
    t["Elements"] = el
    ap = []
    for i, (code, scope, req, default, label, reason) in enumerate(M.APPLIC, 1):
        r = {"applicability_id": f"AP{i:03d}", "element_code": code, **scope_row(scope), "required": REQ[req],
             "default_value": default or "", "label_override": label or "", "reason": reason}
        r["specificity"] = sum(1 for k in ["country", "state", "company", "lob", "product", "object", "module"]
                               if k in scope)
        ap.append(r)
    t["Applicability"] = pd.DataFrame(ap)
    t["PickLists"] = pd.DataFrame([(k, v[0], len(v[1])) for k, v in M.PICKLISTS.items()],
                                  columns=["pick_list", "name", "option_count"])
    opts = []
    for k, (_n, options) in M.PICKLISTS.items():
        for j, (code, label) in enumerate(options, 1):
            scopes = [s for p, o, s in M.OPTION_SCOPES if p == k and o == code]
            where = "; ".join(",".join(f"{a}={b}" for a, b in s.items()) for s in scopes) or "all"
            opts.append({"pick_list": k, "option_code": code, "label": label, "sort": j, "applies_where": where,
                         "effective_from": "2025-01-01", "effective_to": "9999-12-31"})
    t["PickListOptions"] = pd.DataFrame(opts)
    t["Rules"] = pd.DataFrame([{"rule_code": c, "name": n, **scope_row(s), "phase": p, "condition": cond,
                                "action": a, "owner": o} for c, n, s, p, cond, a, o in M.RULES])
    t["Rating"] = pd.DataFrame(M.RATING, columns=["lob", "step", "description", "formula", "factor_table"])
    t["RatingFactors"] = pd.DataFrame([{"factor_table": f, "state": s.get("state", STAR), "keys": k, "value": v}
                                       for f, s, k, v in M.FACTORS])
    t["Sources"] = pd.DataFrame(M.SOURCES, columns=["source_code", "kind", "description"])
    # crosswalk: each synonym is a source field mapped to the canonical element (source attributed by naming style)
    cw = []
    for e in M.ELEMENTS:
        for s in e["synonyms"]:
            src = ("LEGACY" if s.isupper() and ("-" in s or "_" in s) and s.startswith("PM-") else
                   "GW_PC" if s[:1].isupper() and s.isalpha() and not s.isupper() else
                   "LEGACY" if s.isupper() else "HOMEBUILT" if s.islower() or "_" in s else "AMS")
            cw.append({"source_code": src, "source_field": s, "element_code": e["element_code"], "match": "synonym",
                       "status": "Proposed"})
        if e["element_code"] in M.FORM_USE:
            cw.append({"source_code": "ACORD", "source_field": f"ACORD {M.FORM_USE[e['element_code']]}: {e['name']}",
                       "element_code": e["element_code"], "match": "form reference", "status": "Proposed"})
    t["Crosswalk"] = pd.DataFrame(cw)
    return t


# ------------------------------------------------------------------------------------------ resolve: the key in action
def resolve(element_code: str, ctx: dict, applic: pd.DataFrame | None = None) -> dict:
    """Most specific applicability row matching the context wins. ctx has country/state/company/lob/product/object/module."""
    e = next(x for x in M.ELEMENTS if x["element_code"] == element_code)
    if e["lobs"] != STAR and ctx.get("lob") not in e["lobs"].split(","):
        return {"required": REQ["NA"], "default": "", "label": e["name"], "reason": "Not used by this line of business", "rule": ""}
    if ctx.get("module") and ctx["module"] not in e["modules"].split(","):
        return {"required": REQ["NA"], "default": "", "label": e["name"], "reason": "Not used by this module", "rule": ""}
    best, best_s = None, -1
    for code, scope, req, default, label, reason in M.APPLIC:
        if code != element_code:
            continue
        sc = dict(scope)
        if all(ctx.get(k) == v for k, v in sc.items()):
            s = len(sc)
            if s > best_s:
                best, best_s = (req, default, label, reason, sc), s
    if best is None:
        return {"required": REQ[e["base_required"]], "default": "", "label": e["name"], "reason": "Base definition", "rule": ""}
    req, default, label, reason, sc = best
    return {"required": REQ[req], "default": default or "", "label": label or e["name"], "reason": reason,
            "rule": ", ".join(f"{k}={v}" for k, v in sc.items())}


# ------------------------------------------------------------------------------------------ synthetic data
STATES = {"FL": ("Miami", "331"), "TX": ("Houston", "770"), "CA": ("Los Angeles", "900"), "NY": ("Brooklyn", "112"),
          "PA": ("Philadelphia", "191"), "NJ": ("Newark", "071"), "IL": ("Chicago", "606"), "GA": ("Atlanta", "303"),
          "OH": ("Columbus", "432"), "MI": ("Detroit", "482"), "ON": ("Toronto", "M5V")}
NAMES = ["Harbor Point Holdings", "Lone Star Fabrication", "Pacific Crest Foods", "Empire Ridge Builders", "Keystone Plumbing",
         "Garden State Logistics", "Lakeshore Retail Group", "Peachtree Couriers", "Buckeye Supply Co", "Great Lakes Distribution",
         "Maple Leaf Warehousing", "Sunrise Bistro Group", "Bayou Carpentry", "Coastal Freight Lines", "Metro Office Partners",
         "Riverbend Manufacturing", "Gulf Breeze Delivery", "Hudson Valley Transit", "Prairie Haulers", "Motor City Service Co",
         "Blue Ridge Masonry", "Seaside Hospitality"]


def synth(seed: int = 7) -> dict[str, pd.DataFrame]:
    rnd = random.Random(seed)
    pol, party, loc, bldg, cov, cls, stc, veh, drv, prem, clm, cft, bill, ces = ([] for _ in range(14))
    n = {"pol": 0, "loc": 0, "b": 0, "c": 0, "k": 0, "v": 0, "d": 0, "clm": 0}
    names = iter(NAMES)

    def new_policy(lob, state, company, product, country="US"):
        n["pol"] += 1
        pno = f"{lob}{'E' if company == 'SI2-ES' else 'A'}{2026000 + n['pol']:07d}"
        eff = date(2026, 1, 1) + timedelta(days=rnd.randrange(0, 270))
        name = next(names)
        pid = f"PTY{n['pol']:04d}"
        party.append({"party_id": pid, "named_insured_name": name, "legal_entity_type": rnd.choice(["CORP", "LLC"]),
                      "fein": "" if country != "US" else f"XX-XXX{rnd.randrange(1000, 9999)}",
                      "naics_code": {"CP": "531120", "WC": "238350", "GL": "722511", "CA": "484110"}[lob]})
        pol.append({"policy_number": pno, "line_of_business": lob, "writing_company": company, "product_code": product,
                    "country": country, "state": state, "policy_effective_date": eff.isoformat(),
                    "policy_expiration_date": (eff.replace(year=eff.year + 1)).isoformat(), "party_id": pid,
                    "producer_code": f"AG{rnd.randrange(100, 999)}", "commission_rate": 12.5 if lob != "WC" else 8.0,
                    "transaction_type": "NB"})
        bill.append({"billing_account_id": f"BA{n['pol']:04d}", "policy_number": pno,
                     "payment_plan": "AUDIT" if lob == "WC" else rnd.choice(["PIF", "QTR", "MON"])})
        return pno

    def add_loc(pno, state, country="US"):
        n["loc"] += 1
        city, z = STATES[state]
        lid = f"LOC{n['loc']:04d}"
        pc = f"{z}{rnd.randrange(10, 99)}" if country == "US" else f"{z} {rnd.randrange(1, 9)}A{rnd.randrange(1, 9)}"
        loc.append({"location_id": lid, "policy_number": pno, "location_address": f"{rnd.randrange(100, 9999)} Commerce Way, {city}",
                    "state": state, "postal_code": pc, "territory": f"{state}-{rnd.randrange(1, 9):02d}",
                    "flood_zone": rnd.choice(["AE", "X", "VE"]) if state == "FL" else ("X" if country == "US" else "")})
        return lid

    def add_cov(pno, lob, code, **vals):
        n["c"] += 1
        cid = f"COV{n['c']:05d}"
        cov.append({"coverage_id": cid, "policy_number": pno, "line_of_business": lob, "coverage_code": code, **vals})
        return cid

    def add_prem(pno, cid, state, amount, company):
        sl = round(amount * {"TX": 0.0485, "NY": 0.036, "CA": 0.03, "FL": 0.05}.get(state, 0.04), 2) if company == "SI2-ES" else None
        stamp = round(amount * {"TX": 0.0015, "NY": 0.0015, "CA": 0.0018}.get(state, 0), 2) if company == "SI2-ES" and state in ("TX", "NY", "CA") else None
        asl = {"CP": "1.0 / 5.1", "WC": "16.0", "GL": "17.1", "CA": "19.4 / 21.2"}[pno[:2]]
        prem.append({"premium_id": f"PRM{len(prem) + 1:05d}", "policy_number": pno, "coverage_id": cid, "state": state,
                     "written_premium": round(amount, 2), "surplus_lines_tax": sl, "stamping_fee": stamp,
                     "annual_statement_line": asl})

    def add_claim(pno, lob, state, cause, paid, reserve, cid, aww=None):
        n["clm"] += 1
        cno = f"CLM{26000 + n['clm']}"
        clm.append({"claim_number": cno, "policy_number": pno, "line_of_business": lob, "state": state,
                    "loss_date": (date(2026, 3, 1) + timedelta(days=rnd.randrange(0, 180))).isoformat(), "cause_of_loss": cause,
                    "claim_status": "CLOSED" if reserve == 0 else "OPEN", "coverage_id": cid, "injured_worker_aww": aww})
        cft.append({"claim_txn_id": f"CT{len(cft) + 1:05d}", "claim_number": cno, "paid_loss": paid, "case_reserve": reserve})
        if paid + reserve > 250_000:
            ces.append({"cession_id": f"CES{len(ces) + 1:04d}", "treaty_id": "PR-XOL-2026", "claim_number": cno,
                        "ceded_loss": round(paid + reserve - 250_000, 2), "ceded_premium": None})

    # --- Commercial Property
    for state, company, country in [("FL", "SI2-ADM", "US"), ("TX", "SI2-ES", "US"), ("CA", "SI2-ADM", "US"),
                                    ("NY", "SI2-ADM", "US"), ("ON", "SI2-ADM", "CA")]:
        pno = new_policy("CP", state, company, "CP-Std 2025", country)
        for _ in range(rnd.choice([1, 2])):
            lid = add_loc(pno, state, country)
            n["b"] += 1
            bid = f"BLD{n['b']:04d}"
            const = rnd.choice(["1", "2", "4", "6"])
            roof = rnd.choice([2004, 2011, 2016, 2021]) if state == "FL" else rnd.randrange(2002, 2024)
            bldg.append({"building_id": bid, "location_id": lid, "construction_class": const, "year_built": rnd.randrange(1965, 2020),
                         "roof_year": roof, "square_feet": rnd.randrange(8, 120) * 1000, "stories": rnd.choice([1, 2, 3, 6]),
                         "protection_class": str(rnd.randrange(1, 7)), "sprinkler": rnd.choice(["FULL", "NONE", "PART"])})
            blim, bpp = rnd.randrange(20, 150) * 100_000, rnd.randrange(5, 40) * 100_000
            ded_type = "WIND_PCT" if state in ("FL", "TX") else ("EQ_PCT" if state == "CA" else "FLAT")
            cid = add_cov(pno, "CP", "BLDG", building_id=bid, building_limit=blim, bpp_limit=bpp,
                          business_income_limit=rnd.randrange(5, 30) * 100_000, causes_of_loss_form="SPECIAL",
                          valuation="RC", coinsurance="90", deductible_type=ded_type, deductible_amount=10_000,
                          wind_deductible_pct=(2 if state in ("FL", "TX") else None), eq_deductible_pct=(10 if state == "CA" else None))
            rate = 0.71 if state == "FL" else 0.42
            add_prem(pno, cid, state, (blim + bpp) / 100 * rate * (0.7 if bldg[-1]["sprinkler"] == "FULL" else 1) * 1.12, company)
        add_claim(pno, "CP", state, {"FL": "WIND", "TX": "HAIL", "CA": "FIRE", "NY": "WATER", "ON": "FIRE"}[state],
                  rnd.randrange(20, 400) * 1000, rnd.choice([0, 50_000, 180_000]), cid)

    # --- Workers' Compensation
    for state, company, extra in [("TX", "SI2-ADM", None), ("CA", "SI2-ADM", None), ("NY", "SI2-ADM", None),
                                  ("PA", "SI2-ADM", "OH"), ("FL", "SI2-ADM", None)]:
        pno = new_policy("WC", state, company, "WC-Std 2025")
        bureau = {"CA": "WCIRB", "NY": "NYCIRB", "NJ": "NJCRIB", "PA": "PCRB"}.get(state, "NCCI")
        mod = round(rnd.uniform(0.78, 1.62), 2)
        stc.append({"state_coverage_id": f"STC{len(stc) + 1:04d}", "policy_number": pno, "state": state, "wc_part3_type": "3A",
                    "wc_bureau": bureau, "experience_mod": mod})
        if extra:  # monopolistic state on 3.C only: stop-gap, no class payroll
            stc.append({"state_coverage_id": f"STC{len(stc) + 1:04d}", "policy_number": pno, "state": extra, "wc_part3_type": "3C",
                        "wc_bureau": "STATE_FUND", "experience_mod": None})
        cid = add_cov(pno, "WC", "WC-EL", el_limit_each_accident=1_000_000, el_limit_disease_policy=1_000_000,
                      el_limit_disease_employee=1_000_000)
        manual = 0
        for code in rnd.sample(["8810", "5645", "5183", "8742", "9015"], 2):
            n["k"] += 1
            payroll = rnd.randrange(3, 40) * 50_000
            rate = {"8810": 0.2, "8742": 0.35, "5645": 8.5, "5183": 4.1, "9015": 3.2}[code] * (1.9 if state == "NY" else 1)
            cls.append({"class_id": f"CLS{n['k']:04d}", "policy_number": pno, "coverage_id": cid, "lob": "WC", "state": state,
                        "wc_class_code": code, "estimated_payroll": payroll, "full_time_employees": max(1, payroll // 55_000),
                        "gl_class_code": None, "gl_premium_basis": None, "gl_exposure": None, "subcontractor_cost": None,
                        "ny_labor_law_exposure": None})
            manual += payroll / 100 * rate
        add_prem(pno, cid, state, manual * mod + 160, company)
        add_claim(pno, "WC", state, "INJURY", rnd.randrange(5, 90) * 1000, rnd.choice([0, 35_000, 120_000]), cid,
                  aww=rnd.randrange(700, 1600))

    # --- General Liability
    for state, company, glc, basis, product in [("NY", "SI2-ADM", "CONTR-CARP", "PAYROLL", "GL-Contractors 2025"),
                                                ("FL", "SI2-ADM", "RESTAURANT", "SALES", "GL-Std 2025"),
                                                ("TX", "SI2-ES", "CONTR-PLUMB", "PAYROLL", "GL-ClaimsMade 2025"),
                                                ("IL", "SI2-ADM", "RETAIL", "SALES", "GL-Std 2025"),
                                                ("GA", "SI2-ADM", "MFG-LIGHT", "SALES", "GL-Std 2025")]:
        pno = new_policy("GL", state, company, product)
        lid = add_loc(pno, state)
        cm = product == "GL-ClaimsMade 2025"
        cid = add_cov(pno, "GL", "CGL", gl_coverage_form="CM" if cm else "OCC", retroactive_date="2021-07-01" if cm else None,
                      each_occurrence_limit=1_000_000, general_aggregate_limit=2_000_000, products_aggregate_limit=2_000_000,
                      personal_adv_injury_limit=1_000_000, deductible_amount=rnd.choice([0, 2_500, 5_000]))
        n["k"] += 1
        expo = rnd.randrange(4, 60) * 100_000
        contr = glc.startswith("CONTR")
        cls.append({"class_id": f"CLS{n['k']:04d}", "policy_number": pno, "coverage_id": cid, "lob": "GL", "state": state,
                    "location_id": lid, "wc_class_code": None, "estimated_payroll": expo if basis == "PAYROLL" else None,
                    "full_time_employees": None, "gl_class_code": glc, "gl_premium_basis": basis, "gl_exposure": expo,
                    "subcontractor_cost": round(expo * rnd.uniform(0.2, 0.8)) if contr else None,
                    "ny_labor_law_exposure": (True if state == "NY" and contr else None)})
        rate = {"CONTR-CARP": 41.0 if state == "NY" else 18.5, "CONTR-PLUMB": 14.2, "RESTAURANT": 1.95, "RETAIL": 0.9, "MFG-LIGHT": 1.4}[glc]
        add_prem(pno, cid, state, expo / 1000 * rate, company)
        add_claim(pno, "GL", state, rnd.choice(["BI", "PD"]), rnd.randrange(0, 60) * 1000, rnd.choice([25_000, 90_000, 300_000]), cid)

    # --- Commercial Auto
    for state, company, fleet in [("FL", "SI2-ADM", "LT"), ("NY", "SI2-ADM", "PPT"), ("TX", "SI2-ES", "HT"),
                                  ("GA", "SI2-ADM", "MT"), ("MI", "SI2-ADM", "LT")]:
        pno = new_policy("CA", state, company, "CA-Std 2025")
        pip = {"FL": 10_000, "NY": 50_000, "MI": 250_000}.get(state)
        cid = add_cov(pno, "CA", "AUTO-LIAB", auto_symbol="1" if state != "TX" else "7", ca_liability_limit=1_000_000,
                      pip_limit=pip, um_limit=1_000_000 if state != "TX" else 30_000, deductible_amount=1_000)
        prem_total = 0
        for _ in range(rnd.randrange(2, 4)):
            n["v"] += 1
            vt = fleet if rnd.random() < 0.7 else "PPT"
            gvw = {"PPT": 5000, "LT": 9500, "MT": 16000, "HT": 33000, "XHT": 60000, "TRAC": 80000}[vt]
            radius = "LONG" if vt == "HT" and state == "TX" else rnd.choice(["LOCAL", "INTERMEDIATE"])
            veh.append({"vehicle_id": f"VEH{n['v']:04d}", "policy_number": pno, "vin": f"1FT{rnd.randrange(10**13, 10**14)}",
                        "vehicle_year": rnd.randrange(2017, 2026), "vehicle_type": vt, "gvw": gvw, "radius_class": radius,
                        "vehicle_use": "COMMERCIAL" if vt in ("MT", "HT") else "SERVICE",
                        "garaging_zip": STATES[state][1] + f"{rnd.randrange(10, 99)}", "cost_new": rnd.randrange(35, 180) * 1000})
            base = 2150 if state == "FL" else 1450
            prem_total += base * {"PPT": 0.8, "LT": 1.0, "MT": 1.6, "HT": 2.7}[vt] * (1.25 if radius == "LONG" else 1)
            n["d"] += 1
            drv.append({"driver_id": f"DRV{n['d']:04d}", "policy_number": pno, "vehicle_id": f"VEH{n['v']:04d}",
                        "driver_license_state": state, "driver_mvr_points": rnd.choice([0, 0, 1, 2, 3, 6]),
                        "cdl_ind": vt in ("HT", "XHT", "TRAC") and rnd.random() < 0.8})
        add_prem(pno, cid, state, prem_total + ({"FL": 210, "NY": 330}.get(state, 0)), company)
        add_claim(pno, "CA", state, "COLL", rnd.randrange(3, 70) * 1000, rnd.choice([0, 40_000, 400_000]), cid)

    for p in prem:
        if p["written_premium"] > 0:
            ces.append({"cession_id": f"CES{len(ces) + 1:04d}", "treaty_id": "QS-20-2026", "claim_number": None,
                        "ceded_loss": None, "ceded_premium": round(p["written_premium"] * 0.20, 2), "premium_id": p["premium_id"]})
    df = lambda rows: pd.DataFrame(rows)
    out = {"Policy": df(pol), "Party": df(party), "Location": df(loc), "Building": df(bldg), "Coverage": df(cov),
           "Classification": df(cls), "StateCoverage": df(stc), "Vehicle": df(veh), "Driver": df(drv), "Premium": df(prem),
           "Claim": df(clm), "ClaimFinancial": df(cft), "BillingAccount": df(bill), "Cession": df(ces)}
    out["RuleHits"] = rule_hits(out)
    return out


def rule_hits(d: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Run the illustrative business rules on the synthetic data so the demo shows rules firing by key."""
    hits = []
    pol = d["Policy"].set_index("policy_number")
    b = d["Building"].merge(d["Location"], on="location_id")
    for r in b.itertuples():
        st = r.state
        if st == "FL" and r.roof_year < 2026 - 15:
            hits.append(("UW-CP-002", r.policy_number, r.building_id, f"FL roof year {r.roof_year} older than 15 years"))
        elif r.roof_year < 2026 - 20:
            hits.append(("UW-CP-001", r.policy_number, r.building_id, f"Roof year {r.roof_year} older than 20 years"))
        if r.construction_class == "1" and r.stories > 4:
            hits.append(("UW-CP-003", r.policy_number, r.building_id, "Frame building over 4 stories"))
        if st == "FL" and r.flood_zone in ("V", "VE"):
            hits.append(("UW-CP-004", r.policy_number, r.location_id, f"Flood zone {r.flood_zone}"))
    for r in d["StateCoverage"].itertuples():
        if r.state in ("OH", "WA", "ND", "WY") and r.wc_part3_type == "3A":
            hits.append(("UW-WC-001", r.policy_number, r.state_coverage_id, "Monopolistic state listed in 3.A"))
        if r.experience_mod and r.experience_mod > 1.5:
            hits.append(("UW-WC-002", r.policy_number, r.state_coverage_id, f"Experience mod {r.experience_mod}"))
    for r in d["Classification"].query("lob == 'GL'").itertuples():
        if r.state == "NY" and str(r.gl_class_code).startswith("CONTR") and r.ny_labor_law_exposure:
            hits.append(("UW-GL-001", r.policy_number, r.class_id, "NY contractor with Labor Law 240/241 exposure"))
        if r.subcontractor_cost and r.subcontractor_cost / r.gl_exposure > 0.5:
            hits.append(("UW-GL-002", r.policy_number, r.class_id, f"Subcontracted {r.subcontractor_cost / r.gl_exposure:.0%} of exposure"))
    v = d["Vehicle"].merge(d["Driver"], on=["policy_number", "vehicle_id"])
    for r in v.itertuples():
        if r.vehicle_type in ("HT", "XHT", "TRAC") and r.radius_class == "LONG":
            hits.append(("UW-CA-001", r.policy_number, r.vehicle_id, "Long-haul heavy truck"))
        if r.driver_mvr_points >= 6:
            hits.append(("UW-CA-002", r.policy_number, r.driver_id, f"{r.driver_mvr_points} MVR points"))
        if r.vehicle_type in ("HT", "XHT", "TRAC") and not r.cdl_ind:
            hits.append(("UW-CA-003", r.policy_number, r.driver_id, "Heavy truck driver without CDL"))
    for p in d["Premium"].itertuples():
        if pol.at[p.policy_number, "writing_company"] == "SI2-ES":
            hits.append(("UW-ES-001", p.policy_number, p.premium_id, f"Surplus lines tax {p.surplus_lines_tax:,.2f} added"))
    return pd.DataFrame(hits, columns=["rule_code", "policy_number", "record_id", "finding"])


def applicability_matrix() -> pd.DataFrame:
    """One row per element x example context: shows how the same element resolves differently by key."""
    ctxs = [
        ("CP Florida admitted", dict(country="US", state="FL", company="SI2-ADM", lob="CP", product="CP-Std 2025", module="Policy")),
        ("CP Texas E&S", dict(country="US", state="TX", company="SI2-ES", lob="CP", product="CP-Std 2025", module="Policy")),
        ("CP California", dict(country="US", state="CA", company="SI2-ADM", lob="CP", product="CP-Std 2025", module="Policy")),
        ("CP Ontario, Canada", dict(country="CA", state="ON", company="SI2-ADM", lob="CP", product="CP-Std 2025", module="Policy")),
        ("WC New York", dict(country="US", state="NY", company="SI2-ADM", lob="WC", product="WC-Std 2025", module="Policy")),
        ("WC Ohio", dict(country="US", state="OH", company="SI2-ADM", lob="WC", product="WC-Std 2025", module="Policy")),
        ("WC Texas claims", dict(country="US", state="TX", company="SI2-ADM", lob="WC", product="WC-Std 2025", module="Claims")),
        ("GL New York contractor", dict(country="US", state="NY", company="SI2-ADM", lob="GL", product="GL-Contractors 2025", module="Policy")),
        ("GL Texas E&S claims-made", dict(country="US", state="TX", company="SI2-ES", lob="GL", product="GL-ClaimsMade 2025", module="Policy")),
        ("CA Florida", dict(country="US", state="FL", company="SI2-ADM", lob="CA", product="CA-Std 2025", module="Policy")),
        ("CA Georgia", dict(country="US", state="GA", company="SI2-ADM", lob="CA", product="CA-Std 2025", module="Policy")),
    ]
    picked = ["SURPLUS_LINES_TAX", "STAMPING_FEE", "POSTAL_CODE", "STATE", "FEIN", "WIND_DEDUCTIBLE_PCT", "EQ_DEDUCTIBLE_PCT",
              "FLOOD_ZONE", "WC_CLASS_CODE", "WC_BUREAU", "INJURED_WORKER_AWW", "NY_LABOR_LAW_EXPOSURE", "RETROACTIVE_DATE",
              "PIP_LIMIT", "UM_LIMIT"]
    rows = []
    for el in picked:
        for name, ctx in ctxs:
            r = resolve(el, ctx)
            if r["reason"] in ("Not used by this line of business",):
                continue
            rows.append({"element_code": el, "context": name, **{k: ctx.get(k, STAR) for k in ["country", "state", "company", "lob", "product", "module"]},
                         "required": r["required"], "default": r["default"], "label": r["label"], "matched_scope": r["rule"] or "(base)",
                         "reason": r["reason"]})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------------------------------ SQL Server DDL
TYPE_SQL = {"TEXT": "NVARCHAR(200)", "CODE": "NVARCHAR(40)", "DATE": "DATE", "AMOUNT": "DECIMAL(18,2)", "PERCENT": "DECIMAL(9,4)",
            "DECIMAL": "DECIMAL(9,4)", "INTEGER": "INT", "BOOLEAN": "BIT"}


def sql_lit(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NA:
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return repr(float(v)) if isinstance(v, float) else str(v)
    return "N'" + str(v).replace("'", "''") + "'"


def sql_type_for(col: str, s: pd.Series) -> str:
    el = next((e for e in M.ELEMENTS if e["element_code"].lower() == col), None)
    vals = s.dropna()
    if el and el["data_type"] not in ("TEXT", "CODE"):
        return TYPE_SQL[el["data_type"]]
    if len(vals) and vals.map(lambda x: isinstance(x, bool)).all():
        return "BIT"
    if pd.api.types.is_float_dtype(s):
        return "DECIMAL(18,4)"
    if pd.api.types.is_integer_dtype(s):
        return "INT"
    longest = int(vals.astype(str).str.len().max()) if len(vals) else 0
    for size in (40, 120, 400, 1000):
        if longest <= size * 0.8:
            return f"NVARCHAR({size})"
    return "NVARCHAR(MAX)"


def sql_script(meta: dict[str, pd.DataFrame], data: dict[str, pd.DataFrame]) -> str:
    out = ["-- SI² Commercial Lines CDM (Commercial Property, Workers' Compensation, General Liability, Commercial Auto)",
           f"-- Version {VERSION}. Generated from the same store as SI2_Commercial_Ontology.ttl. Illustrative codes and synthetic data.",
           "-- Run in your database (for example SI2Canonical). Re-running drops and recreates schema cdm objects.",
           "SET NOCOUNT ON;", "GO",
           "IF NOT EXISTS (SELECT 1 FROM sys.schemas WHERE name = 'cdm') EXEC('CREATE SCHEMA cdm');", "GO"]

    def table(name, df, pk=None, fks=()):
        cols = list(df.columns)
        out.append(f"IF OBJECT_ID('cdm.{name}') IS NOT NULL DROP TABLE cdm.[{name}];")
        defs = []
        for c in cols:
            t = sql_type_for(c, df[c])
            null = "NOT NULL" if c == pk else "NULL"
            defs.append(f"  [{c}] {t} {null}")
        if pk:
            defs.append(f"  CONSTRAINT PK_{name} PRIMARY KEY ([{pk}])")
        out.append(f"CREATE TABLE cdm.[{name}] (\n" + ",\n".join(defs) + "\n);")
        rows = [", ".join(sql_lit(None if (not isinstance(v, (list, dict)) and pd.isna(v)) else v) for v in r)
                for r in df.astype(object).itertuples(index=False)]
        for i in range(0, len(rows), 500):
            out.append(f"INSERT INTO cdm.[{name}] ({', '.join(f'[{c}]' for c in cols)}) VALUES\n(" + "),\n(".join(rows[i:i + 500]) + ");")
        out.append("GO")

    out.append("-- ===== Metadata: the governed store (one row per object, element, scope, option, rule)")
    order = list(data)[::-1]  # drop children first is not needed (no FK constraints declared across reloads)
    for name, pk in [("Keys", "key"), ("Objects", "object"), ("Relations", None), ("Elements", "element_code"),
                     ("Applicability", "applicability_id"), ("PickLists", "pick_list"), ("PickListOptions", None),
                     ("Rules", "rule_code"), ("Rating", None), ("RatingFactors", None), ("Sources", "source_code"),
                     ("Crosswalk", None)]:
        df = meta[name].copy()
        df.columns = [c.replace(" ", "_") for c in df.columns]
        table("Meta" + name, df, pk.replace(" ", "_") if pk else None)
    out.append("-- ===== Entities: the physical CDM with synthetic data for CP, WC, GL and CA")
    pks = {o[0]: o[4] for o in M.OBJECTS}
    for name in ["Party", "Policy", "BillingAccount", "Location", "Building", "Coverage", "Classification", "StateCoverage",
                 "Vehicle", "Driver", "Premium", "Claim", "ClaimFinancial", "Cession", "RuleHits"]:
        table(name, data[name], pks.get(name))
    out.append("""-- ===== Resolve an element for a context: most specific matching scope wins
CREATE OR ALTER FUNCTION cdm.fn_ResolveElement (@element NVARCHAR(40), @country NVARCHAR(10), @state NVARCHAR(10),
       @company NVARCHAR(20), @lob NVARCHAR(10), @product NVARCHAR(60), @module NVARCHAR(20), @asof DATE)
RETURNS TABLE AS RETURN
-- No row returned = use cdm.MetaElements.base_required.
SELECT TOP 1 a.element_code, a.required, a.default_value, a.label_override, a.reason, a.applicability_id
FROM cdm.MetaApplicability a
WHERE a.element_code = @element
  AND a.country IN ('*', @country) AND a.state IN ('*', @state) AND a.company IN ('*', @company)
  AND a.lob IN ('*', @lob) AND a.product IN ('*', @product) AND a.module IN ('*', @module)
  AND @asof BETWEEN a.effective_from AND a.effective_to
ORDER BY a.specificity DESC;
GO
-- Example: SELECT * FROM cdm.fn_ResolveElement('PIP_LIMIT','US','FL','SI2-ADM','CA','CA-Std 2025','Policy','2026-06-01');
CREATE OR ALTER VIEW cdm.vw_PolicyPremium AS
SELECT p.policy_number, p.line_of_business, p.writing_company, p.state, pa.named_insured_name,
       SUM(pr.written_premium) AS written_premium, SUM(pr.surplus_lines_tax) AS surplus_lines_tax
FROM cdm.Policy p JOIN cdm.Party pa ON pa.party_id = p.party_id JOIN cdm.Premium pr ON pr.policy_number = p.policy_number
GROUP BY p.policy_number, p.line_of_business, p.writing_company, p.state, pa.named_insured_name;
GO""")
    return "\n".join(out) + "\n"


# ------------------------------------------------------------------------------------------ ontology
SI2 = Namespace("https://si2insure.com/ontology/commercial#")
XSD_FOR = {"TEXT": XSD.string, "CODE": XSD.string, "DATE": XSD.date, "AMOUNT": XSD.decimal, "PERCENT": XSD.decimal,
           "DECIMAL": XSD.decimal, "INTEGER": XSD.integer, "BOOLEAN": XSD.boolean}


def ontology(meta: dict[str, pd.DataFrame]) -> Graph:
    g = Graph()
    g.bind("si2", SI2); g.bind("skos", SKOS); g.bind("owl", OWL); g.bind("dcterms", DCTERMS)
    ont = URIRef(str(SI2).rstrip("#"))
    g.add((ont, RDF.type, OWL.Ontology))
    g.add((ont, DCTERMS.title, Literal("SI² Commercial Lines Ontology (CP, WC, GL, CA)")))
    g.add((ont, OWL.versionInfo, Literal(VERSION)))
    g.add((ont, RDFS.comment, Literal("Generated from the same governed store as the SI² Commercial CDM. Codes are illustrative.")))
    # meta classes
    for c, label in [("ContextKey", "Context key"), ("Applicability", "Applicability (element in a context)"),
                     ("BusinessRule", "Business rule"), ("RatingStep", "Rating step"), ("SourceSystem", "Source system"),
                     ("LineOfBusiness", "Line of business"), ("Module", "Module"), ("WritingCompany", "Writing company")]:
        g.add((SI2[c], RDF.type, OWL.Class)); g.add((SI2[c], RDFS.label, Literal(label)))
    for p, label in [("appliesToElement", "applies to element"), ("required", "required"), ("defaultValue", "default value"),
                     ("labelOverride", "label override"), ("reason", "reason"), ("effectiveFrom", "effective from"),
                     ("effectiveTo", "effective to"), ("specificity", "specificity"), ("condition", "condition"),
                     ("action", "action"), ("phase", "phase"), ("formula", "formula"), ("stepNumber", "step"),
                     ("usesElement", "uses element"), ("usedInModule", "used in module"), ("usedInLob", "used in line of business"),
                     ("hasPickList", "has pick list"), ("acordForm", "ACORD form"), ("mapsFrom", "maps from source field")]:
        g.add((SI2[p], RDF.type, OWL.AnnotationProperty if p in ("reason", "acordForm") else RDF.Property))
        g.add((SI2[p], RDFS.label, Literal(label)))
    for k in M.KEY:
        g.add((SI2["key_" + k], RDF.type, OWL.DatatypeProperty)); g.add((SI2["key_" + k], RDFS.label, Literal(k)))
        g.add((SI2["key_" + k], RDFS.subPropertyOf, SI2.contextKey))
        g.add((SI2["key_" + k], RDFS.domain, SI2.Applicability))
    for r in meta["Keys"].itertuples():
        for k in str(r.key).split(" / "):
            g.add((SI2["key_" + k], SKOS.definition, Literal(f"{r.meaning}. {r._4}")))
    for code, name in M.LOBS.items():
        g.add((SI2["LOB_" + code], RDF.type, SI2.LineOfBusiness)); g.add((SI2["LOB_" + code], RDFS.label, Literal(name)))
    for code, name in M.COMPANIES.items():
        g.add((SI2["CO_" + code.replace("-", "_")], RDF.type, SI2.WritingCompany)); g.add((SI2["CO_" + code.replace("-", "_")], RDFS.label, Literal(name)))
    for m in ["Policy", "Billing", "Claims", "Reinsurance", "Finance", "Reporting"]:
        g.add((SI2["MOD_" + m], RDF.type, SI2.Module)); g.add((SI2["MOD_" + m], RDFS.label, Literal(m)))
    # objects -> classes
    for code, name, parent, definition, key in M.OBJECTS:
        g.add((SI2[code], RDF.type, OWL.Class)); g.add((SI2[code], RDFS.label, Literal(name)))
        g.add((SI2[code], SKOS.definition, Literal(definition)))
        if parent:
            g.add((SI2[code], RDFS.subClassOf, SI2[parent]))
    for s, rel, o, card in M.RELATIONS:
        g.add((SI2[rel], RDF.type, OWL.ObjectProperty)); g.add((SI2[rel], RDFS.label, Literal(rel)))
        bn = SI2[f"{s}_{rel}_{o}"]
        g.add((bn, RDF.type, OWL.ObjectProperty)); g.add((bn, RDFS.subPropertyOf, SI2[rel]))
        g.add((bn, RDFS.domain, SI2[s])); g.add((bn, RDFS.range, SI2[o])); g.add((bn, RDFS.comment, Literal(f"cardinality {card}")))
    # pick lists -> SKOS concept schemes
    for k, (name, options) in M.PICKLISTS.items():
        sch = SI2["PL_" + k]
        g.add((sch, RDF.type, SKOS.ConceptScheme)); g.add((sch, RDFS.label, Literal(name)))
        for code, label in options:
            c = SI2[f"PL_{k}_{code.replace('-', '_')}"]
            g.add((c, RDF.type, SKOS.Concept)); g.add((c, SKOS.inScheme, sch)); g.add((c, SKOS.notation, Literal(code)))
            g.add((c, SKOS.prefLabel, Literal(label)))
    for i, (k, code, scope) in enumerate(M.OPTION_SCOPES, 1):
        c = SI2[f"PL_{k}_{code.replace('-', '_')}"]
        g.add((c, SKOS.scopeNote, Literal("applies where " + ", ".join(f"{a}={b}" for a, b in scope.items()))))
    # elements -> datatype properties
    for e in M.ELEMENTS:
        p = SI2[e["element_code"]]
        g.add((p, RDF.type, OWL.DatatypeProperty)); g.add((p, RDFS.label, Literal(e["name"])))
        g.add((p, SKOS.definition, Literal(e["definition"]))); g.add((p, RDFS.domain, SI2[e["object"]]))
        g.add((p, RDFS.range, XSD_FOR[e["data_type"]]))
        for s in e["synonyms"]:
            g.add((p, SKOS.altLabel, Literal(s)))
        if e["pick_list"]:
            g.add((p, SI2.hasPickList, SI2["PL_" + e["pick_list"]]))
        for lob in (M.LOBS if e["lobs"] == STAR else e["lobs"].split(",")):
            g.add((p, SI2.usedInLob, SI2["LOB_" + lob]))
        for m in e["modules"].split(","):
            g.add((p, SI2.usedInModule, SI2["MOD_" + m]))
        if e["element_code"] in M.FORM_USE:
            g.add((p, SI2.acordForm, Literal(M.FORM_USE[e["element_code"]])))
    # applicability -> individuals carrying the context key
    for r in meta["Applicability"].itertuples(index=False):
        a = SI2[r.applicability_id]
        g.add((a, RDF.type, SI2.Applicability)); g.add((a, SI2.appliesToElement, SI2[r.element_code]))
        for k in M.KEY:
            g.add((a, SI2["key_" + k], Literal(getattr(r, k))))
        g.add((a, SI2.required, Literal(r.required))); g.add((a, SI2.specificity, Literal(int(r.specificity))))
        if r.default_value:
            g.add((a, SI2.defaultValue, Literal(r.default_value)))
        if r.label_override:
            g.add((a, SI2.labelOverride, Literal(r.label_override)))
        g.add((a, SI2.reason, Literal(r.reason)))
    # rules and rating
    el_codes = [e["element_code"] for e in M.ELEMENTS]
    for r in meta["Rules"].itertuples(index=False):
        u = SI2["RULE_" + r.rule_code.replace("-", "_")]
        g.add((u, RDF.type, SI2.BusinessRule)); g.add((u, RDFS.label, Literal(r.name)))
        g.add((u, SI2.phase, Literal(r.phase))); g.add((u, SI2.condition, Literal(r.condition))); g.add((u, SI2.action, Literal(r.action)))
        for k in M.KEY:
            g.add((u, SI2["key_" + k], Literal(getattr(r, k))))
        for c in el_codes:
            if c in r.condition or c in r.action:
                g.add((u, SI2.usesElement, SI2[c]))
    for lob, step, desc, formula, ft in M.RATING:
        u = SI2[f"RATE_{lob}_{step}"]
        g.add((u, RDF.type, SI2.RatingStep)); g.add((u, RDFS.label, Literal(f"{M.LOBS[lob]} step {step}: {desc}")))
        g.add((u, SI2.stepNumber, Literal(step))); g.add((u, SI2.formula, Literal(formula)))
        g.add((u, SI2.usedInLob, SI2["LOB_" + lob]))
        for c in el_codes:
            if c in formula:
                g.add((u, SI2.usesElement, SI2[c]))
    for code, kind, desc in M.SOURCES:
        g.add((SI2["SRC_" + code], RDF.type, SI2.SourceSystem)); g.add((SI2["SRC_" + code], RDFS.label, Literal(f"{code} ({kind})")))
        g.add((SI2["SRC_" + code], RDFS.comment, Literal(desc)))
    for r in meta["Crosswalk"].itertuples(index=False):
        g.add((SI2[r.element_code], SI2.mapsFrom, Literal(f"{r.source_code}:{r.source_field}")))
    return g


# ------------------------------------------------------------------------------------------ Excel
def write_xlsx(path: Path, meta, data, matrix):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    readme = pd.DataFrame({"SI² Commercial Lines CDM": [
        f"Version {VERSION}. One governed store generates this CDM workbook, the SQL Server script and the ontology (Turtle / JSON-LD).",
        "Lines: Commercial Property (CP), Workers' Compensation (WC), General Liability (GL), Commercial Auto (CA).",
        "Context key: country · state · company · line · LOB · product · object · module · effective dates. '*' = all. Most specific match wins.",
        "Sheets: Keys, Objects, Relations, Elements, Applicability, ResolvedExamples, PickLists, PickListOptions, Rules, Rating, RatingFactors, Sources, Crosswalk, then synthetic data (Policy ... RuleHits).",
        "All codes, rates and data are illustrative and synthetic. Bureau class codes and forms belong to ISO, NCCI, ACORD and the state bureaus; verify against licensed manuals.",
        "Governance: SI² proposes elements and mappings; SMEs approve by key; any change to a source system raises an alert (mandatory) or waits for release approval.",
    ]})
    sheets = {"README": readme, **{k: meta[k] for k in ["Keys", "Objects", "Relations", "Elements", "Applicability"]},
              "ResolvedExamples": matrix, **{k: meta[k] for k in ["PickLists", "PickListOptions", "Rules", "Rating", "RatingFactors",
                                                                   "Sources", "Crosswalk"]}, **data}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0B1B2B")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for i, col in enumerate(df.columns, 1):
                width = min(60, max(10, int(max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) * 1.05)))
                ws.column_dimensions[get_column_letter(i)].width = width if name != "README" else 140
            if len(df):
                ws.auto_filter.ref = ws.dimensions


def main(out: Path):
    out.mkdir(parents=True, exist_ok=True)
    meta, data, matrix = tables(), synth(), applicability_matrix()
    write_xlsx(out / "SI2_Commercial_CDM.xlsx", meta, data, matrix)
    (out / "SI2_Commercial_CDM_SQLServer.sql").write_text(sql_script(meta, data), encoding="utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, df in {**meta, "ResolvedExamples": matrix}.items():
            z.writestr(f"metadata/{name}.csv", df.to_csv(index=False))
        for name, df in data.items():
            z.writestr(f"sample_data/{name}.csv", df.to_csv(index=False))
    (out / "SI2_Commercial_CDM_csv.zip").write_bytes(buf.getvalue())
    g = ontology(meta)
    (out / "SI2_Commercial_Ontology.ttl").write_text(g.serialize(format="turtle"), encoding="utf-8")
    (out / "SI2_Commercial_Ontology.jsonld").write_text(g.serialize(format="json-ld", indent=1), encoding="utf-8")
    files = ["SI2_Commercial_CDM.xlsx", "SI2_Commercial_CDM_SQLServer.sql", "SI2_Commercial_CDM_csv.zip",
             "SI2_Commercial_Ontology.ttl", "SI2_Commercial_Ontology.jsonld"]
    with zipfile.ZipFile(out / "SI2_Commercial_CDM_and_Ontology.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for f in files:
            z.write(out / f, f)
    stats = {"objects": len(M.OBJECTS), "elements": len(M.ELEMENTS), "applicability": len(M.APPLIC),
             "pick_lists": len(M.PICKLISTS), "options": len(meta["PickListOptions"]), "rules": len(M.RULES),
             "rating_steps": len(M.RATING), "triples": len(g), "rule_hits": len(data["RuleHits"]),
             **{f"rows_{k}": len(v) for k, v in data.items()}}
    return stats, meta, data, matrix


if __name__ == "__main__":
    s, *_ = main(Path(sys.argv[1] if len(sys.argv) > 1 else "out"))
    for k, v in s.items():
        print(f"{k:22s} {v}")
