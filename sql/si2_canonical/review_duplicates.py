"""Duplicate review for the SI² Canonical Master: one entry, one meaning.

    python review_duplicates.py                                   write duplicate_review_<time>.xlsx (nothing changes)
    python review_duplicates.py --apply duplicate_review_<time>.xlsx --approve-as "Saba"

The review workbook has three sheets:
  Merge          groups of data elements that look like the same thing, with a suggested element to keep. Rows marked
                 MERGE are applied: the other element's code, name and synonyms become synonyms of the one kept, the
                 kept element takes its pick list or definition when it has none, and the other element is expired
                 (never deleted). Change MERGE to KEEP BOTH, or leave it empty, for anything you do not want merged.
  Bureau codes   ISO / NCCI codes that look like placeholders (ISO_CT1, NCCI_1120). Rows marked EXPIRE are expired.
  Shared names   one name that points to several elements (for information: it is where meanings collide).
Every change is a change request made by 'SI2 review' and approved by the second person.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import difflib
import json
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
REVIEWER = "SI2 review"
ABBR = {"no": "number", "nbr": "number", "num": "number", "amt": "amount", "dt": "date", "cd": "code", "ind": "indicator",
        "flg": "indicator", "flag": "indicator", "pct": "percent", "qty": "quantity", "addr": "address", "desc": "description",
        "id": "identifier", "cov": "coverage", "ded": "deductible", "prem": "premium", "pol": "policy", "clm": "claim",
        "eff": "effective", "exp": "expiration", "yr": "year", "veh": "vehicle", "drv": "driver", "dob": "birthdate",
        "tel": "phone", "telephone": "phone", "postal": "zip", "zipcode": "zip", "nm": "name", "dte": "date"}
STOP = {"the", "of", "a", "an", "and"}           # 'by', 'on', 'to' change the meaning (Reviewed By / Reviewed On)
WEAK = {"code", "value", "indicator"}           # words that rarely change the meaning: Claim Status = Claim Status Code
TEMPLATE = re.compile(r"^(business data element representing|configured value or category|date or time associated with|"
                      r"descriptive information documenting|current status or state of|unique identifier for|standardized|"
                      r"human-readable|recorded|set of)", re.I)


def norm(s) -> str:
    return re.sub(r"[ _\-./]", "", str(s or "")).upper()


def tokens(text) -> tuple:
    t = str(text or "")
    if "." in t and " " not in t:
        t = t.split(".")[-1]
    t = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", t)
    t = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1 \2", t)
    out = []
    for w in re.split(r"[^A-Za-z0-9]+", t.lower()):
        if not w or w in STOP:
            continue
        w = ABBR.get(w, w)
        if len(w) > 3 and w.endswith("s") and not w.endswith("ss"):
            w = w[:-1]
        out.append(w)
    return tuple(out)


class UnionFind:
    def __init__(self):
        self.p = {}

    def find(self, x):
        self.p.setdefault(x, x)
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        self.p[self.find(a)] = self.find(b)


def find_duplicates(elements: list[dict], synonyms: list[dict], rule_use: dict) -> tuple[list[dict], list[dict]]:
    """elements: {id, code, name, desc, dtype, picklist, bureau, created_by}; synonyms: {id, text, source}.
    Returns (merge rows, shared-name rows)."""
    by_id = {e["id"]: e for e in elements}
    edges = defaultdict(list)                     # (a, b) -> evidence

    def link(a, b, strength, why):
        if a != b:
            edges[tuple(sorted((a, b)))].append((strength, why))

    # 1. the same words in the name (any order, abbreviations expanded)
    groups = defaultdict(list)
    for e in elements:
        tk = frozenset(tokens(e["name"])) or frozenset(tokens(e["code"]))
        if tk:
            groups[tk].append(e["id"])
    for tk, ids in groups.items():
        ids = sorted(set(ids))
        if 1 < len(ids) <= 12:
            for i in ids[1:]:
                a, b = by_id[ids[0]], by_id[i]
                full = frozenset(tokens(a["code"])) == frozenset(tokens(b["code"]))
                link(ids[0], i, 2 if full else 1, ("same words: " if full else "same words in the name: ") + " ".join(sorted(tk)))
    # 2. the same words apart from 'code' / 'value' / 'indicator' (needs two meaningful words)
    core = defaultdict(list)
    for e in elements:
        c = frozenset(tokens(e["name"])) - WEAK
        if len(c) >= 2:
            core[c].append(e["id"])
    for c, ids in core.items():
        ids = sorted(set(ids))
        if 1 < len(ids) <= 8:
            for i in ids[1:]:
                if frozenset(tokens(by_id[ids[0]]["name"])) != frozenset(tokens(by_id[i]["name"])):
                    link(ids[0], i, 1, "same words apart from " + "/".join(sorted(WEAK)))
    # 3. names: one element's own code or name is another element's synonym (strong); or both carry the same synonym (weaker)
    own = defaultdict(set)
    for e in elements:
        for t in (e["code"], e["name"]):
            own[norm(t)].add(e["id"])
    syn_owner = defaultdict(set)
    for s in synonyms:
        if s["id"] in by_id and norm(s["text"]):
            syn_owner[norm(s["text"])].add(s["id"])
    for key, holders in syn_owner.items():
        for i in own.get(key, ()):
            for h in holders:
                if h != i:
                    link(h, i, 2, f"'{key}' is the name of one and a synonym of the other")
        if 1 < len(holders) <= 3:
            hs = sorted(holders)
            for i in hs[1:]:
                link(hs[0], i, 1, f"both have the synonym '{key}'")
    shared = []
    for key in set(own) | set(syn_owner):
        ids = own.get(key, set()) | syn_owner.get(key, set())
        if len(ids) > 1 and key:
            shared.append({"name": key, "elements": ", ".join(sorted(by_id[i]["code"] for i in ids)), "count": len(ids)})
    # 4. the same (non-template) business definition, when the names share a word
    defs = defaultdict(list)
    for e in elements:
        d = re.sub(r"\W+", " ", str(e.get("desc") or "")).strip().lower()
        if len(d) >= 30 and not TEMPLATE.match(d):
            defs[d].append(e["id"])
    for d, ids in defs.items():
        if 1 < len(ids) <= 6:
            for i in ids[1:]:
                ta, tb = set(tokens(by_id[ids[0]]["name"])) - WEAK, set(tokens(by_id[i]["name"])) - WEAK
                if ta and tb and len(ta & tb) / len(ta | tb) >= 0.34:
                    link(ids[0], i, 1, "same business definition")

    uf = UnionFind()
    for (a, b), ev in edges.items():
        if by_id[a]["dtype"] != by_id[b]["dtype"] and max(s for s, _ in ev) < 2:
            continue                              # weak evidence across data types: leave apart
        uf.union(a, b)
    clusters = defaultdict(list)
    for x in list(uf.p):
        clusters[uf.find(x)].append(x)

    def keep_score(e):
        return (100 if e.get("created_by") == "SI2 seed" else 0) + min(rule_use.get(e["id"], 0), 5) * 10 + \
               (10 if e.get("picklist") else 0) + (10 if e.get("bureau") else 0) + \
               (0 if TEMPLATE.match(str(e.get("desc") or "")) else min(len(str(e.get("desc") or "")) // 40, 8)) - len(e["code"]) / 20 - e["id"] / 1e6 \
               - (15 if re.search(r"_(OPTIONS?|LIST|VALUES?)$", e["code"]) or e["code"].startswith(("GLOBAL_", "DE_")) else 0)

    rows, n = [], 0
    for members in sorted(clusters.values(), key=lambda m: min(by_id[i]["code"] for i in m)):
        if len(members) < 2:
            continue
        n += 1
        keep = max((by_id[i] for i in members), key=keep_score)
        for i in sorted(members, key=lambda i: by_id[i]["code"]):
            if i == keep["id"]:
                continue
            e = by_id[i]
            ev = edges.get(tuple(sorted((i, keep["id"]))), [])
            strong = any(s == 2 for s, _ in ev)
            evidence = "; ".join(sorted({w for _, w in ev})) or "same group (through another member)"
            sim = difflib.SequenceMatcher(None, str(e.get("desc") or "").lower(), str(keep.get("desc") or "").lower()).ratio()
            ta, tb = set(tokens(e["name"])) - WEAK - {"option"}, set(tokens(keep["name"])) - WEAK - {"option"}
            jac = len(ta & tb) / len(ta | tb) if ta | tb else 0
            specific = len(ta) >= 2 and len(tb) >= 2
            words = any(s == 2 and w.startswith("same words") for s, w in ev)
            count_vs_id = (e["code"].startswith("NUMBER_") or e["code"].startswith("NUM_")) != \
                          (keep["code"].startswith("NUMBER_") or keep["code"].startswith("NUM_"))
            confidence = "High" if (strong and specific and e["dtype"] == keep["dtype"] and len(members) <= 6
                                    and (words or jac >= 0.5) and not count_vs_id) else "Medium"
            rows.append({"Group": n, "Confidence": confidence, "Evidence": evidence,
                         "Keep": keep["code"], "Keep name": keep["name"], "Keep type": keep["dtype"], "Keep definition": keep.get("desc"),
                         "Merge": e["code"], "Merge name": e["name"], "Merge type": e["dtype"], "Merge definition": e.get("desc"),
                         "Definition similarity": round(sim, 2), "Rules using the merged element": rule_use.get(i, 0),
                         "Decision": "MERGE" if confidence == "High" else "", "Note": ""})
    shared.sort(key=lambda r: (-r["count"], r["name"]))
    return rows, shared


# ------------------------------------------------------------------------------------------ database
def connect(server, database):
    import pyodbc
    drivers = ["ODBC Driver 18 for SQL Server", "ODBC Driver 17 for SQL Server", "SQL Server Native Client 11.0", "SQL Server"]
    driver = next((d for d in drivers if d in pyodbc.drivers()), None)
    if not driver:
        sys.exit("No SQL Server ODBC driver found.")
    return pyodbc.connect(f"DRIVER={{{driver}}};SERVER={server};DATABASE={database};Trusted_Connection=yes;TrustServerCertificate=yes;",
                          autocommit=True)


def q(con, sql, *p):
    cur = con.cursor()
    cur.execute(sql, *p)
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def load(con):
    elements = q(con, "SELECT e.ElementId AS id, e.ElementCode AS code, e.Name AS name, e.BusinessDescription AS [desc], e.DataType AS dtype, "
                      "p.PickListCode AS picklist, e.IsBureauReported AS bureau, e.CreatedBy AS created_by FROM canon.DataElement e "
                      "LEFT JOIN canon.PickList p ON p.PickListId = e.PickListId WHERE e.ScopeId = 1 AND e.Status = 'ACTIVE'")
    synonyms = q(con, "SELECT ObjectId AS id, SynonymText AS text, SourceSystem AS source FROM canon.Synonym "
                      "WHERE ObjectType = 'ELEMENT' AND Status = 'ACTIVE' AND ScopeId = 1")
    use = defaultdict(int)
    for r in q(con, "SELECT ElementId AS id, COUNT(*) AS n FROM rules.RuleCondition GROUP BY ElementId "
                    "UNION ALL SELECT TargetElementId, COUNT(*) FROM rules.BusinessRule WHERE TargetElementId IS NOT NULL GROUP BY TargetElementId"):
        use[r["id"]] += r["n"]
    bureau = q(con, "SELECT b.BureauCodeMapId AS id, p.PickListCode AS picklist, o.OptionCode AS [option], o.OptionName AS option_name, "
                    "b.Bureau AS bureau, b.BureauCode AS code, b.CreatedBy AS created_by FROM canon.BureauCodeMap b "
                    "JOIN canon.PickListOption o ON o.OptionId = b.OptionId JOIN canon.PickList p ON p.PickListId = o.PickListId "
                    "WHERE b.Status = 'ACTIVE' AND (b.BureauCode LIKE 'ISO[_]%' OR b.BureauCode LIKE 'NCCI[_]%' OR b.BureauCode LIKE 'ISO-%')")
    return elements, synonyms, dict(use), bureau


# ------------------------------------------------------------------------------------------ workbook
SHEETS = ("Merge", "Bureau codes", "Shared names")


def write_review(path: Path, merges, bureau, shared) -> Path:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        out = path.with_suffix(".csv")
        with open(out, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(merges[0]) if merges else ["Group"])
            w.writeheader()
            w.writerows(merges)
        return out
    wb = openpyxl.Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3A93")
    for title, rows, choices in ((SHEETS[0], merges, '"MERGE,KEEP BOTH"'), (SHEETS[1], bureau, '"EXPIRE,KEEP"'), (SHEETS[2], shared, None)):
        ws = wb.active if title == SHEETS[0] else wb.create_sheet(title)
        ws.title = title
        if not rows:
            ws.append(["Nothing found."])
            continue
        cols = list(rows[0])
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head, fill, Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([r.get(c) for c in cols])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = {"Evidence": 34, "Keep definition": 50, "Merge definition": 50, "Keep name": 26, "Merge name": 26, "Keep": 28, "Merge": 28,
                  "Note": 30, "elements": 80, "name": 30}
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(c, 14)
        if choices and "Decision" in cols:
            col = openpyxl.utils.get_column_letter(cols.index("Decision") + 1)
            dv = DataValidation(type="list", formula1=choices, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(rows) + 1}")
    wb.save(path)
    return path


def read_review(path: Path) -> tuple[list[dict], list[dict]]:
    if path.suffix.lower() == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as fh:
            return list(csv.DictReader(fh)), []
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)

    def sheet(name):
        if name not in wb.sheetnames:
            return []
        rows = list(wb[name].iter_rows(values_only=True))
        if not rows or rows[0][0] == "Nothing found.":
            return []
        return [dict(zip(rows[0], r)) for r in rows[1:]]
    return sheet(SHEETS[0]), sheet(SHEETS[1])


def plan_apply(merges: list[dict], bureau: list[dict], elements: list[dict], synonyms: list[dict]) -> list[dict]:
    by_code = {e["code"]: e for e in elements}
    known = defaultdict(set)
    for e in elements:
        known[e["code"]] |= {norm(e["code"]), norm(e["name"])}
    by_id = {e["id"]: e for e in elements}
    syn_of = defaultdict(list)
    for s in synonyms:
        if s["id"] in by_id:
            syn_of[by_id[s["id"]]["code"]].append(s)
            known[by_id[s["id"]]["code"]].add(norm(s["text"]))
    out, merged = [], set()
    today = dt.date.today().isoformat()
    for r in merges:
        if str(r.get("Decision") or "").strip().upper() != "MERGE":
            continue
        keep, gone = str(r.get("Keep") or "").strip(), str(r.get("Merge") or "").strip()
        if keep not in by_code or gone not in by_code or keep == gone or gone in merged or keep in merged:
            continue
        merged.add(gone)
        k, g = by_code[keep], by_code[gone]
        texts = [(g["code"], "*", "BUSINESS"), (g["name"], "*", "BUSINESS")] + [(s["text"], s["source"], "SOURCE") for s in syn_of[gone]]
        for text, src, typ in texts:
            if text and norm(text) not in known[keep] and len(str(text)) <= 300:
                known[keep].add(norm(text))
                out.append({"object_type": "SYNONYM", "change": "ADD", "code": keep, "summary": f"Merge: '{text}' now means {keep}",
                            "payload": {"object_type": "ELEMENT", "object_code": keep, "text": text, "type": typ, "source": (src or "*")[:60],
                                        "method": "Duplicate review"}})
        upd = {}
        if not k.get("picklist") and g.get("picklist"):
            upd.update({"picklist": g["picklist"], "data_type": "CODE"})
        if (not k.get("desc") or TEMPLATE.match(str(k.get("desc")))) and g.get("desc") and not TEMPLATE.match(str(g.get("desc"))):
            upd["description"] = str(g["desc"])[:2000]
        if upd:
            out.append({"object_type": "ELEMENT", "change": "UPDATE", "code": keep, "summary": f"Merge: {keep} takes {', '.join(upd)} from {gone}",
                        "payload": {"code": keep, **upd}})
        out.append({"object_type": "ELEMENT", "change": "EXPIRE", "code": gone, "summary": f"Merge: {gone} expired, now a synonym of {keep}",
                    "payload": {"code": gone, "effective_to": today}})
    for r in bureau:
        if str(r.get("Decision") or "").strip().upper() == "EXPIRE" and r.get("id"):
            out.append({"object_type": "BUREAU_CODE", "change": "EXPIRE", "code": f"{r.get('picklist')}.{r.get('option')}.{r.get('bureau')}",
                        "summary": f"Placeholder {r.get('bureau')} code {r.get('code')} expired", "payload": {"id": int(r["id"]), "effective_to": today}})
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--server", default=os.environ.get("SI2_MSSQL_SERVER", "SABAHOME\\SQLEXPRESS02"))
    ap.add_argument("--database", default="SI2Canonical")
    ap.add_argument("--apply", default=None, help="the reviewed workbook to apply ('latest' = newest duplicate_review_*.xlsx)")
    ap.add_argument("--approve-as", default=None)
    a = ap.parse_args()
    con = connect(a.server, a.database)
    elements, synonyms, use, bureau = load(con)
    if not a.apply:
        merges, shared = find_duplicates(elements, synonyms, use)
        brows = [{**b, "Decision": "EXPIRE", "Note": ""} for b in bureau]
        path = write_review(HERE / f"duplicate_review_{dt.datetime.now():%Y%m%d_%H%M}.xlsx", merges, brows, shared)
        groups = len({r["Group"] for r in merges})
        print(f"Active data elements: {len(elements)}")
        print(f"Likely duplicates: {len(merges)} element(s) in {groups} group(s) - High {sum(r['Confidence'] == 'High' for r in merges)}, "
              f"Medium {sum(r['Confidence'] == 'Medium' for r in merges)} (High rows are pre-marked MERGE)")
        print(f"Placeholder-looking bureau codes: {len(brows)} (pre-marked EXPIRE)")
        print(f"Names shared by several elements: {len(shared)}")
        print(f"Review: {path}")
        return 0
    path = Path(a.apply)
    if a.apply == "latest":
        files = sorted(list(HERE.glob("duplicate_review_*.xlsx")) + list(HERE.glob("duplicate_review_*.csv")), key=lambda f: f.stem)
        if not files:
            sys.exit("No duplicate_review_*.xlsx found.")
        path = files[-1]
    merges, brows = read_review(path)
    changes = plan_apply(merges, brows, elements, synonyms)
    print(f"{path.name}: {sum(1 for c in changes if c['change'] == 'EXPIRE' and c['object_type'] == 'ELEMENT')} merge(s), "
          f"{sum(1 for c in changes if c['object_type'] == 'SYNONYM')} synonym(s) moved, "
          f"{sum(1 for c in changes if c['object_type'] == 'BUREAU_CODE')} bureau code(s) expired")
    if not changes:
        return 0
    cur = con.cursor()
    cur.execute("SET NOCOUNT ON; IF NOT EXISTS (SELECT 1 FROM intake.SourceSystem WHERE SourceSystemCode = N'Duplicate review') "
                "INSERT intake.SourceSystem (SourceSystemCode, Name, Kind) VALUES (N'Duplicate review', N'SI2 duplicate review', 'APP'); "
                "INSERT intake.ScanRun (SourceSystemCode, DatabaseName, StartedBy, Message) VALUES (N'Duplicate review', DB_NAME(), ?, ?); "
                "SELECT CAST(SCOPE_IDENTITY() AS int);", REVIEWER, path.name)
    run = cur.fetchone()[0]
    for c in changes:
        cur.execute("EXEC gov.usp_RequestChange @ObjectType = ?, @ChangeType = ?, @ObjectCode = ?, @PayloadJson = ?, @RequestedBy = ?, "
                    "@Summary = ?, @Source = 'INTAKE', @ScanRunId = ?", c["object_type"], c["change"], c["code"][:300],
                    json.dumps(c["payload"], ensure_ascii=False), REVIEWER, c["summary"][:1000], run)
    print(f"Proposed as '{REVIEWER}' (review run {run}).")
    if a.approve_as:
        cur.execute("EXEC gov.usp_DecideAll @DecidedBy = ?, @Approve = 1, @Note = N'Duplicate review', @ScanRunId = ?", a.approve_as, run)
        sets = []
        while True:
            if cur.description:
                sets.append(cur.fetchall())
            if not cur.nextset():
                break
        print("Decided: " + ", ".join(f"{r[0]} {r[1]}" for r in (sets[0] if sets else [])))
        for r in (sets[1] if len(sets) > 1 else [])[:20]:
            print(f"  FAILED {r[1]} {r[3]}: {r[5]}")
    c2 = con.cursor()
    c2.execute("SELECT * FROM canon.vw_Summary")
    row = c2.fetchone()
    print("Master now:", ", ".join(f"{d[0]} {v}" for d, v in zip(c2.description, row)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
